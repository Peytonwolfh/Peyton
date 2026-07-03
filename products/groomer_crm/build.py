#!/usr/bin/env python3
"""
Dog Groomer Client CRM — Etsy digital product builder (re-skin of the
Pet Sitter CRM generator; see products/petsitter_crm/build.py).

Generates Groomer-Client-CRM.xlsx (Google Sheets + Excel compatible):
  * standard formulas only (SUMIFS, COUNTIFS, IF, IFERROR, VLOOKUP, INDEX/MATCH,
    LARGE, TODAY, DATE, WEEKDAY, TEXT) — no VBA, no macros, no Excel-only functions
  * data-validation dropdowns, formula-based conditional formatting
  * parameterized CONFIG so the same script can be re-skinned for other verticals
    by editing CONFIG only.

Run:  python3 build.py          -> builds + verifies the workbook
"""

import os
import re
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))

# ============================================================================
# CONFIG — everything vertical-specific lives here.
# ============================================================================

CONFIG = {
    "output_file": "Groomer-Client-CRM.xlsx",
    "product_name": "Dog Groomer Client CRM",
    "tagline": "Client, pet, cut-history, appointment & payment tracker for dog groomers",

    # ---- theme: warm dusty rose / mauve + charcoal -------------------------
    "theme": {
        "charcoal":      "413A3E",  # header rows (warm charcoal)
        "charcoal_dark": "322C30",  # banners
        "accent_dark":   "9D6B7B",  # dusty rose — accents / section heads
        "accent_mid":    "C49AA6",  # tile labels / tab color
        "accent_light":  "F2E4E8",  # tile bodies / today-highlight
        "accent_pale":   "F8F1F3",  # page background feel
        "cream":         "FBF9F7",  # input cells
        "calc_fill":     "F2EEEF",  # auto-calculated cells (do not type)
        "red_fill":      "F6D9D5",
        "red_text":      "9C2B23",
        "amber_fill":    "FBEFD0",
        "amber_text":    "8A6116",
        "green_fill":    "DDEBD8",
        "green_text":    "3E6B47",
        "sample_text":   "6D6D6D",  # sample rows: gray italic
        "white":         "FFFFFF",
    },

    # ---- sheet display names (order = tab order) -------------------------
    "sheet_names": {
        "start":    "START HERE",
        "dash":     "Dashboard",
        "clients":  "Clients",
        "pets":     "Pets",
        "visits":   "Appointments",
        "invoices": "Invoices & Payments",
        "services": "Services & Rates",
    },

    # ---- vertical vocabulary ----------------------------------------------
    "vocab": {
        "client": "Client",
        "unit": "Dog",
        "unit_plural": "Dogs",
        "visit": "Appointment",
        "caution_keyword": "CAUTION",
    },

    # ---- dropdown option lists (inline lists, comma-safe values) ---------
    "lists": {
        "client_status": ["Active", "Inactive", "Waitlist"],
        "contact_method": ["Text", "Call", "Email"],
        "coat_type": ["Double Coat", "Curly", "Wire", "Silky", "Smooth"],
        "visit_status": ["Booked", "Done", "Cancelled"],
        "paid": ["Yes", "No"],
        "invoice_status": ["Draft", "Sent", "Paid", "Overdue"],
    },

    # ---- editable price list (Services & Rates tab) -----------------------
    # These are NOT sample rows — they're the real starter price list.
    "services": [
        ("Full Groom — Small",  65.00, "90 minutes",  "Under 25 lbs. Bath, dry, haircut, nails, ears"),
        ("Full Groom — Medium", 85.00, "2 hours",     "25–60 lbs. Bath, dry, haircut, nails, ears"),
        ("Full Groom — Large", 110.00, "2.5–3 hours", "60+ lbs. Doodles / heavy coats may add a fee"),
        ("Bath & Brush",        45.00, "60 minutes",  "Bath, blow-dry, brush-out, nails, ears — no haircut"),
        ("Nail Trim",           15.00, "15 minutes",  "Walk-in friendly; +$5 with Dremel grinding"),
        ("De-shed Treatment",   60.00, "90 minutes",  "Double coats: undercoat rake + high-velocity dry"),
        ("Puppy Intro Groom",   40.00, "45 minutes",  "Under 6 months. Face/feet/tail tidy + handling practice"),
    ],

    # ---- column definitions ------------------------------------------------
    # type: input | dropdown:<listkey> | dropdown_range:<defined name> |
    #       date | date_calc | currency_calc | currency_input
    "columns": {
        "clients": [
            ("Client Name",               22, "input"),
            ("Phone",                     15, "input"),
            ("Email",                     24, "input"),
            ("Preferred Contact",         16, "dropdown:contact_method"),
            ("Emergency Contact (name & phone)", 26, "input"),
            ("Vet Clinic",                22, "input"),
            ("Vet Phone",                 15, "input"),
            ("Status",                    12, "dropdown:client_status"),
            ("Notes",                     32, "input"),
        ],
        "pets": [
            ("Owner",                     22, "dropdown_range:ClientList"),
            ("Dog Name",                  14, "input"),
            ("Breed",                     18, "input"),
            ("Age",                        6, "input"),
            ("Coat Type",                 14, "dropdown:coat_type"),
            ("Last Cut / Style Notes (what the owner asked for)", 38, "input"),
            ("Blade / Comb Lengths",      22, "input"),
            ("Shampoo & Products Used",   26, "input"),
            ("Skin Conditions / Allergies", 26, "input"),
            ("Behavior Notes / Warnings", 32, "input"),
            ("Matting Notes",             26, "input"),
            ("Ear / Nail / Gland Prefs",  28, "input"),
        ],
        "visits": [
            ("Date",                      12, "date"),
            ("Client",                    22, "dropdown_range:ClientList"),
            ("Dog",                       14, "input"),
            ("Service",                   20, "dropdown_range:ServiceList"),
            ("Rate ($)",                  10, "currency_calc"),
            ("Status",                    11, "dropdown:visit_status"),
            ("Paid?",                      8, "dropdown:paid"),
            ("Rebook (weeks)",            14, "input"),
            ("Next Groom Due",            15, "date_calc"),
            ("Notes",                     30, "input"),
        ],
        "invoices": [
            ("Invoice #",                 11, "input"),
            ("Client",                    22, "dropdown_range:ClientList"),
            ("Period Start",              13, "date"),
            ("Period End",                13, "date"),
            ("Invoice Total ($)",         15, "currency_calc"),
            ("Balance Due ($)",           15, "currency_calc"),
            ("Date Sent",                 12, "date"),
            ("Date Paid",                 12, "date"),
            ("Status",                    11, "dropdown:invoice_status"),
            ("Notes",                     30, "input"),
        ],
        "services": [
            ("Service",                   22, "input"),
            ("Rate ($)",                  10, "currency_input"),
            ("Duration",                  12, "input"),
            ("Notes",                     44, "input"),
        ],
    },

    # ---- how deep formulas / dropdowns are pre-loaded ----------------------
    "rows": {
        "clients_max": 200,    # ClientList range + dropdown depth
        "services_max": 30,    # ServiceList range
        "pets_max": 300,
        "visits_max": 250,     # rate formula pre-filled to here
        "invoices_max": 100,   # SUMIFS pre-filled to here
        "top_clients": 5,      # dashboard top-N
        "top_engine": 20,      # first N client rows scanned for top clients
    },

    # ---- sample rows (gray italic, "delete me" note, listed in START HERE) --
    # Dates are TODAY()-relative formulas so the demo conditional formatting
    # (today's appointments, unpaid+done, rebook-due) actually lights up.
    "sample": {
        "clients": [
            ["Amanda Foster", "(555) 214-7788", "amanda.f@example.com",
             "Text", "Mike Foster (555) 214-7790",
             "Lakeview Animal Hospital", "(555) 640-2210",
             "Active", "SAMPLE ROW — delete me"],
            ["Derek Okafor", "(555) 318-2276", "derek.o@example.com",
             "Call", "Nia Okafor (555) 318-9901",
             "Eastside Vet Clinic", "(555) 772-0184",
             "Active", "Prefers early morning slots — SAMPLE ROW, delete me"],
            ["Priya Raman", "(555) 467-9032", "priya.r@example.com",
             "Email", "Dev Raman (555) 467-9033",
             "Harbor Animal Hospital", "(555) 903-3341",
             "Waitlist", "SAMPLE ROW — delete me"],
        ],
        "pets": [
            ["Amanda Foster", "Murphy", "Goldendoodle", "3", "Curly",
             "Teddy-bear head, 1/2\" body all over — owner wants the ears "
             "left long this time",
             "#4F body, #10 sanitary, 1/2\" comb legs",
             "Oatmeal shampoo + light conditioner",
             "Sensitive skin — no perfumed finishing sprays",
             "Sweet but wiggly on the table — go slow on feet",
             "Mats behind ears & armpits when owner skips brushing — check "
             "every visit; charge de-mat fee over 15 min",
             "Grind nails (hates clippers); pluck ears; no glands (vet does) "
             "— SAMPLE ROW, delete me"],
            ["Derek Okafor", "Luna", "Siberian Husky", "5", "Double Coat",
             "De-shed only — NEVER shave (owner confirmed twice)",
             "No blades — undercoat rake + high-velocity dry",
             "De-shedding shampoo + conditioner",
             "None",
             "Vocal but friendly — sings through the whole dry",
             "None — coat well maintained between visits",
             "Nails clipped OK; ears wipe only — SAMPLE ROW, delete me"],
            ["Priya Raman", "Taco", "Chihuahua", "8", "Smooth",
             "Bath & brush; tidy face only, no body clipping",
             "#10 on face if needed",
             "Hypoallergenic shampoo",
             "Allergic to chicken-based treats",
             "CAUTION: nippy when paws are handled — muzzle for nails, "
             "treats after, keep sessions short",
             "None",
             "Nail trim every visit; glands yes — SAMPLE ROW, delete me"],
            ["Amanda Foster", "Bella", "Shih Tzu", "6", "Silky",
             "Short puppy cut, rounded face, bow on collar to finish",
             "#7F body, 3/8\" comb legs",
             "Whitening shampoo (white coat), tearless face wash",
             "Tear staining — clean around eyes gently",
             "Calm, easy on the table",
             "Fine mats in leg feathering if she goes past 6 weeks",
             "Nails clipped; ears plucked; glands yes — SAMPLE ROW, delete me"],
        ],
        # date offset (days from TODAY, or a literal formula string),
        # client, dog, service, status, paid, rebook_weeks, note
        "visits": [
            [0,   "Amanda Foster", "Murphy", "Full Groom — Medium",
             "Booked", "No",  6, "SAMPLE — today's appointments glow"],
            [-3,  "Priya Raman", "Taco", "Nail Trim",
             "Done",   "No",  4, "SAMPLE — done + unpaid turns red"],
            [-7,  "Derek Okafor", "Luna", "De-shed Treatment",
             "Done",   "No",  8, "SAMPLE ROW — delete me"],
            # anchored to the 1st of the current month so the "Revenue This
            # Month" dashboard tile always has something to show
            ["=DATE(YEAR(TODAY()),MONTH(TODAY()),1)",
             "Amanda Foster", "Bella", "Full Groom — Small",
             "Done",   "Yes", 6, "SAMPLE ROW — delete me"],
            [-35, "Amanda Foster", "Bella", "Full Groom — Small",
             "Done",   "Yes", 4, "SAMPLE — overdue rebook turns amber"],
            [2,   "Derek Okafor", "Luna", "Bath & Brush",
             "Booked", "No",  None, "SAMPLE ROW — delete me"],
        ],
        # inv#, client, period_start_offset, period_end_offset,
        # sent_offset (None = blank), paid_offset (None = blank), status, note
        "invoices": [
            ["INV-1001", "Amanda Foster", -40, -30, -28, -25, "Paid",
             "SAMPLE ROW — delete me"],
            ["INV-1002", "Derek Okafor", -8, 0, 0, None, "Sent",
             "SAMPLE ROW — delete me"],
            ["INV-1003", "Priya Raman", -20, -2, -18, None, "Overdue",
             "SAMPLE ROW — delete me"],
        ],
    },
}

T = CONFIG["theme"]
S = CONFIG["sheet_names"]
R = CONFIG["rows"]

# ============================================================================
# style helpers
# ============================================================================

def fill(hex_):
    return PatternFill("solid", fgColor=hex_)

THIN = Side(style="thin", color="E0D6D9")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_FONT = Font(name="Calibri", bold=True, color=T["white"], size=11)
HEADER_FILL = fill(T["charcoal"])
SAMPLE_FONT = Font(name="Calibri", italic=True, color=T["sample_text"], size=11)
BODY_FONT = Font(name="Calibri", size=11)

CURRENCY_FMT = '"$"#,##0.00'
DATE_FMT = "mm/dd/yyyy"


def style_header_row(ws, n_cols, row=1, height=26):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[row].height = height


def set_widths(ws, cols):
    for i, (_, width, _) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def apply_col_formats(ws, cols, first_row, last_row):
    """Number formats + calc-cell shading for whole data region."""
    for i, (_, _, ctype) in enumerate(cols, start=1):
        for r in range(first_row, last_row + 1):
            cell = ws.cell(row=r, column=i)
            if ctype in ("currency_calc", "currency_input"):
                cell.number_format = CURRENCY_FMT
            elif ctype in ("date", "date_calc"):
                cell.number_format = DATE_FMT
            if ctype in ("currency_calc", "date_calc"):
                cell.fill = fill(T["calc_fill"])


def mark_sample_rows(ws, n_cols, first_row, n_rows):
    for r in range(first_row, first_row + n_rows):
        for c in range(1, n_cols + 1):
            ws.cell(row=r, column=c).font = SAMPLE_FONT


def add_dv(ws, dv_type, formula1, cell_range):
    dv = DataValidation(type=dv_type, formula1=formula1, allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv


def inline_list(key):
    return '"' + ",".join(CONFIG["lists"][key]) + '"'


def col_letter(sheet_key, label):
    """Column letter for a labeled column on a data sheet."""
    labels = [l for l, _, _ in CONFIG["columns"][sheet_key]]
    return get_column_letter(labels.index(label) + 1)


def build_columns(ws, sheet_key, max_row):
    """Headers, widths, dropdown DV for a standard data sheet."""
    cols = CONFIG["columns"][sheet_key]
    for i, (label, _, _) in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=label)
    style_header_row(ws, len(cols))
    set_widths(ws, cols)
    ws.freeze_panes = "A2"
    for i, (_, _, ctype) in enumerate(cols, start=1):
        col = get_column_letter(i)
        rng = f"{col}2:{col}{max_row}"
        if ctype.startswith("dropdown_range:"):
            add_dv(ws, "list", ctype.split(":", 1)[1], rng)
        elif ctype.startswith("dropdown:"):
            add_dv(ws, "list", inline_list(ctype.split(":", 1)[1]), rng)
    apply_col_formats(ws, cols, 2, max_row)
    return cols


# ============================================================================
# sheet builders
# ============================================================================

def build_services(wb):
    ws = wb[S["services"]]
    cols = build_columns(ws, "services", R["services_max"])
    for r, row in enumerate(CONFIG["services"], start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    # gentle banding on the real price rows
    for r in range(2, len(CONFIG["services"]) + 2):
        for c in range(1, len(cols) + 1):
            ws.cell(row=r, column=c).border = BORDER_ALL
    note = ws.cell(row=R["services_max"] + 2, column=1,
                   value="Edit names and rates freely — the Appointments tab "
                         "looks prices up from this list automatically. "
                         "Service names here must match the dropdown picks "
                         "exactly (they will, if you always pick from the "
                         "dropdown).")
    note.font = Font(italic=True, color=T["sample_text"], size=10)
    ws.merge_cells(start_row=R["services_max"] + 2, start_column=1,
                   end_row=R["services_max"] + 2, end_column=4)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_properties.tabColor = T["accent_dark"]


def build_clients(wb):
    ws = wb[S["clients"]]
    build_columns(ws, "clients", R["clients_max"])
    samples = CONFIG["sample"]["clients"]
    for r, row in enumerate(samples, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    mark_sample_rows(ws, len(CONFIG["columns"]["clients"]), 2, len(samples))
    # status column CF: Active = green tint, Waitlist = amber
    status_col = col_letter("clients", "Status")
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{R['clients_max']}",
        FormulaRule(formula=[f'${status_col}2="Active"'],
                    fill=fill(T["green_fill"]),
                    font=Font(color=T["green_text"], bold=True)))
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{R['clients_max']}",
        FormulaRule(formula=[f'${status_col}2="Waitlist"'],
                    fill=fill(T["amber_fill"]),
                    font=Font(color=T["amber_text"], bold=True)))
    ws.sheet_properties.tabColor = T["accent_mid"]


def build_pets(wb):
    ws = wb[S["pets"]]
    cols = build_columns(ws, "pets", R["pets_max"])
    samples = CONFIG["sample"]["pets"]
    for r, row in enumerate(samples, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    # CAUTION keyword anywhere in the behavior column -> whole row red
    beh_col = col_letter("pets", "Behavior Notes / Warnings")
    last_col = get_column_letter(len(cols))
    kw = CONFIG["vocab"]["caution_keyword"]
    ws.conditional_formatting.add(
        f"A2:{last_col}{R['pets_max']}",
        FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{kw}",${beh_col}2))'],
            fill=fill(T["red_fill"]),
            font=Font(color=T["red_text"]),
            stopIfTrue=True))
    ws.sheet_properties.tabColor = T["accent_mid"]


def build_visits(wb):
    ws = wb[S["visits"]]
    cols = build_columns(ws, "visits", R["visits_max"])
    labels = [l for l, _, _ in cols]
    c_date = col_letter("visits", "Date")                # A
    c_service = col_letter("visits", "Service")          # D
    c_status = col_letter("visits", "Status")            # F
    c_paid = col_letter("visits", "Paid?")               # G
    c_rebook = col_letter("visits", "Rebook (weeks)")    # H
    c_due = col_letter("visits", "Next Groom Due")       # I
    last_col = get_column_letter(len(cols))
    svc_sheet = S["services"]

    # pre-load the rate lookup + rebook due-date all the way down
    for r in range(2, R["visits_max"] + 1):
        ws.cell(row=r, column=labels.index("Rate ($)") + 1).value = (
            f'=IF(${c_service}{r}="","",'
            f"IFERROR(VLOOKUP(${c_service}{r},"
            f"'{svc_sheet}'!$A$2:$B${R['services_max']},2,FALSE),\"\"))"
        )
        ws.cell(row=r, column=labels.index("Next Groom Due") + 1).value = (
            f'=IF(OR(${c_date}{r}="",${c_rebook}{r}=""),"",'
            f"${c_date}{r}+${c_rebook}{r}*7)"
        )

    samples = CONFIG["sample"]["visits"]
    for r, (off, client, dog, service, status, paid, rebook, note) in \
            enumerate(samples, start=2):
        if isinstance(off, str):
            date_formula = off
        else:
            date_formula = (f"=TODAY()+{off}" if off >= 0
                            else f"=TODAY()-{abs(off)}")
        ws.cell(row=r, column=1, value=date_formula).number_format = DATE_FMT
        ws.cell(row=r, column=2, value=client)
        ws.cell(row=r, column=3, value=dog)
        ws.cell(row=r, column=4, value=service)
        ws.cell(row=r, column=6, value=status)
        ws.cell(row=r, column=7, value=paid)
        if rebook is not None:
            ws.cell(row=r, column=8, value=rebook)
        ws.cell(row=r, column=10, value=note)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    # keep number formats after sample write
    apply_col_formats(ws, cols, 2, R["visits_max"])

    # CF 1 (wins): done + not paid -> red row
    ws.conditional_formatting.add(
        f"A2:{last_col}{R['visits_max']}",
        FormulaRule(
            formula=[f'AND(${c_status}2="Done",${c_paid}2<>"Yes")'],
            fill=fill(T["red_fill"]), font=Font(color=T["red_text"]),
            stopIfTrue=True))
    # CF 2: today's appointments -> dusty rose highlight
    ws.conditional_formatting.add(
        f"A2:{last_col}{R['visits_max']}",
        FormulaRule(formula=[f"${c_date}2=TODAY()"],
                    fill=fill(T["accent_light"]),
                    font=Font(color=T["charcoal_dark"], bold=True),
                    stopIfTrue=True))
    # CF 3: rebook due date has arrived -> amber on the Next Groom Due cell
    ws.conditional_formatting.add(
        f"{c_due}2:{c_due}{R['visits_max']}",
        FormulaRule(
            formula=[f'AND(ISNUMBER(${c_due}2),${c_due}2<=TODAY())'],
            fill=fill(T["amber_fill"]),
            font=Font(color=T["amber_text"], bold=True)))
    ws.sheet_properties.tabColor = T["accent_mid"]


def build_invoices(wb):
    ws = wb[S["invoices"]]
    cols = build_columns(ws, "invoices", R["invoices_max"])
    labels = [l for l, _, _ in cols]
    c_client = col_letter("invoices", "Client")
    c_start = col_letter("invoices", "Period Start")
    c_end = col_letter("invoices", "Period End")
    c_status = col_letter("invoices", "Status")
    last_col = get_column_letter(len(cols))
    vs = S["visits"]
    v_rate = col_letter("visits", "Rate ($)")
    v_client = col_letter("visits", "Client")
    v_date = col_letter("visits", "Date")
    v_status = col_letter("visits", "Status")
    v_paid = col_letter("visits", "Paid?")

    for r in range(2, R["invoices_max"] + 1):
        common = (f"'{vs}'!${v_rate}:${v_rate},"
                  f"'{vs}'!${v_client}:${v_client},${c_client}{r},"
                  f"'{vs}'!${v_status}:${v_status},\"Done\","
                  f"'{vs}'!${v_date}:${v_date},\">=\"&${c_start}{r},"
                  f"'{vs}'!${v_date}:${v_date},\"<=\"&${c_end}{r}")
        guard = (f'OR(${c_client}{r}="",${c_start}{r}="",${c_end}{r}="")')
        ws.cell(row=r, column=labels.index("Invoice Total ($)") + 1).value = \
            f'=IF({guard},"",SUMIFS({common}))'
        ws.cell(row=r, column=labels.index("Balance Due ($)") + 1).value = \
            (f'=IF({guard},"",SUMIFS({common},'
             f"'{vs}'!${v_paid}:${v_paid},\"<>Yes\"))")

    def off_formula(off):
        if off is None:
            return None
        return f"=TODAY()+{off}" if off >= 0 else f"=TODAY()-{abs(off)}"

    samples = CONFIG["sample"]["invoices"]
    for r, (num, client, ps, pe, sent, paid, status, note) in \
            enumerate(samples, start=2):
        ws.cell(row=r, column=1, value=num)
        ws.cell(row=r, column=2, value=client)
        ws.cell(row=r, column=3, value=off_formula(ps))
        ws.cell(row=r, column=4, value=off_formula(pe))
        if sent is not None:
            ws.cell(row=r, column=7, value=off_formula(sent))
        if paid is not None:
            ws.cell(row=r, column=8, value=off_formula(paid))
        ws.cell(row=r, column=9, value=status)
        ws.cell(row=r, column=10, value=note)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    apply_col_formats(ws, cols, 2, R["invoices_max"])

    # status CF on whole row
    rng = f"A2:{last_col}{R['invoices_max']}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'${c_status}2="Overdue"'], fill=fill(T["red_fill"]),
        font=Font(color=T["red_text"]), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'${c_status}2="Paid"'], fill=fill(T["green_fill"]),
        font=Font(color=T["green_text"]), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'${c_status}2="Sent"'], fill=fill(T["amber_fill"]),
        font=Font(color=T["amber_text"]), stopIfTrue=True))
    ws.sheet_properties.tabColor = T["accent_mid"]


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

def _tile(ws, col, label, formula, num_fmt, label_row, value_row):
    """Two merged 2-wide cells stacked: label on top, big value below."""
    c1, c2 = col, get_column_letter(column_index_from_string(col) + 1)
    ws.merge_cells(f"{c1}{label_row}:{c2}{label_row}")
    ws.merge_cells(f"{c1}{value_row}:{c2}{value_row}")
    lab = ws[f"{c1}{label_row}"]
    lab.value = label.upper()
    lab.font = Font(bold=True, size=9, color=T["white"])
    lab.fill = fill(T["accent_dark"])
    lab.alignment = Alignment(horizontal="center", vertical="center")
    val = ws[f"{c1}{value_row}"]
    val.value = formula
    val.font = Font(bold=True, size=22, color=T["charcoal_dark"])
    val.fill = fill(T["accent_light"])
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.number_format = num_fmt
    for cc in (c1, c2):
        ws[f"{cc}{label_row}"].fill = fill(T["accent_dark"])
        ws[f"{cc}{value_row}"].fill = fill(T["accent_light"])


def build_dashboard(wb):
    ws = wb[S["dash"]]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = T["charcoal_dark"]

    cl, vs = S["clients"], S["visits"]
    cl_status = col_letter("clients", "Status")
    v_date = col_letter("visits", "Date")
    v_client = col_letter("visits", "Client")
    v_rate = col_letter("visits", "Rate ($)")
    v_status = col_letter("visits", "Status")
    v_paid = col_letter("visits", "Paid?")

    # widths: A spacer, B..I tiles, then helper cols
    ws.column_dimensions["A"].width = 2
    for c in "BCEFHI":
        ws.column_dimensions[c].width = 13
    for c in "DG":
        ws.column_dimensions[c].width = 2
    ws.column_dimensions["J"].width = 2
    ws.column_dimensions["O"].width = 12
    ws.column_dimensions["P"].width = 18

    # title band
    ws.merge_cells("B2:I2")
    t = ws["B2"]
    t.value = CONFIG["product_name"].upper() + "  —  DASHBOARD"
    t.font = Font(bold=True, size=16, color=T["white"])
    t.fill = fill(T["charcoal_dark"])
    t.alignment = Alignment(horizontal="center", vertical="center")
    for c in "CDEFGHI":
        ws[f"{c}2"].fill = fill(T["charcoal_dark"])
    ws.row_dimensions[2].height = 30

    ws["B3"] = ('="Everything on this page updates itself — nothing to type '
                'here.  Today: "&TEXT(TODAY(),"mmm d, yyyy")')
    ws["B3"].font = Font(italic=True, size=10, color=T["sample_text"])

    week_start = "(TODAY()-WEEKDAY(TODAY(),2)+1)"
    month_start = "DATE(YEAR(TODAY()),MONTH(TODAY()),1)"
    next_month = "DATE(YEAR(TODAY()),MONTH(TODAY())+1,1)"
    tiles = [
        ("B", "Active Clients",
         f"=COUNTIF('{cl}'!${cl_status}:${cl_status},\"Active\")", "0"),
        ("E", "Dogs Groomed This Month",
         f"=COUNTIFS('{vs}'!${v_status}:${v_status},\"Done\","
         f"'{vs}'!${v_date}:${v_date},\">=\"&{month_start},"
         f"'{vs}'!${v_date}:${v_date},\"<\"&{next_month})", "0"),
        ("H", "Appointments This Week",
         f"=COUNTIFS('{vs}'!${v_date}:${v_date},\">=\"&{week_start},"
         f"'{vs}'!${v_date}:${v_date},\"<=\"&{week_start}+6,"
         f"'{vs}'!${v_status}:${v_status},\"<>Cancelled\")", "0"),
    ]
    tiles2 = [
        ("B", "Appointments Today",
         f"=COUNTIFS('{vs}'!${v_date}:${v_date},TODAY(),"
         f"'{vs}'!${v_status}:${v_status},\"<>Cancelled\")", "0"),
        ("E", "Revenue This Month",
         f"=SUMIFS('{vs}'!${v_rate}:${v_rate},"
         f"'{vs}'!${v_status}:${v_status},\"Done\","
         f"'{vs}'!${v_date}:${v_date},\">=\"&{month_start},"
         f"'{vs}'!${v_date}:${v_date},\"<\"&{next_month})",
         CURRENCY_FMT),
        ("H", "Unpaid Balance",
         f"=SUMIFS('{vs}'!${v_rate}:${v_rate},"
         f"'{vs}'!${v_status}:${v_status},\"Done\","
         f"'{vs}'!${v_paid}:${v_paid},\"<>Yes\")",
         CURRENCY_FMT),
    ]
    for col, label, formula, fmt in tiles:
        _tile(ws, col, label, formula, fmt, 5, 6)
    for col, label, formula, fmt in tiles2:
        _tile(ws, col, label, formula, fmt, 8, 9)
    ws.row_dimensions[6].height = 40
    ws.row_dimensions[9].height = 40

    # ---- top clients table -------------------------------------------------
    ws.merge_cells("B12:I12")
    h = ws["B12"]
    h.value = "TOP CLIENTS  (completed appointments, all time)"
    h.font = Font(bold=True, size=12, color=T["white"])
    h.fill = fill(T["charcoal"])
    h.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in "CDEFGHI":
        ws[f"{c}12"].fill = fill(T["charcoal"])
    ws.row_dimensions[12].height = 22

    hdr = {"B": "Rank", "C": "Client", "H": "Revenue"}
    ws.merge_cells("C13:G13")
    for c, v in hdr.items():
        cell = ws[f"{c}13"]
        cell.value = v
        cell.font = Font(bold=True, size=10, color=T["accent_dark"])
        cell.fill = fill(T["accent_pale"])
    ws.merge_cells("H13:I13")
    ws["H13"].fill = fill(T["accent_pale"])

    eng_top, eng_n = 14, R["top_engine"]           # engine rows 14..33
    eng_rng = f"$O${eng_top}:$O${eng_top + eng_n - 1}"
    name_rng = f"$P${eng_top}:$P${eng_top + eng_n - 1}"
    for k in range(1, R["top_clients"] + 1):
        row = 13 + k
        ws[f"B{row}"] = k
        ws[f"B{row}"].alignment = Alignment(horizontal="center")
        ws.merge_cells(f"C{row}:G{row}")
        ws[f"C{row}"] = (f"=IFERROR(INDEX({name_rng},"
                         f"MATCH(LARGE({eng_rng},{k}),{eng_rng},0)),\"—\")")
        ws.merge_cells(f"H{row}:I{row}")
        ws[f"H{row}"] = f'=IFERROR(ROUND(LARGE({eng_rng},{k}),2),"")'
        ws[f"H{row}"].number_format = CURRENCY_FMT
        for c in "BCDEFGHI":
            ws[f"{c}{row}"].fill = fill(
                T["accent_pale"] if k % 2 else T["white"])
        ws[f"C{row}"].font = Font(size=11, color=T["charcoal_dark"])
        ws[f"H{row}"].font = Font(size=11, bold=True, color=T["accent_dark"])

    # ---- hidden-ish calculation engine (cols O:P) ---------------------------
    ws["O13"] = "engine ↓"
    ws["P13"] = "do not edit"
    for c in ("O13", "P13"):
        ws[c].font = Font(size=8, color=T["sample_text"], italic=True)
    for i in range(eng_n):
        r = eng_top + i
        cl_row = 2 + i
        ws[f"O{r}"] = (f"=IF('{cl}'!A{cl_row}=\"\",\"\","
                       f"SUMIFS('{vs}'!${v_rate}:${v_rate},"
                       f"'{vs}'!${v_client}:${v_client},'{cl}'!A{cl_row},"
                       f"'{vs}'!${v_status}:${v_status},\"Done\")"
                       f"+ROW()/10000000)")
        ws[f"P{r}"] = f"=IF($O{r}=\"\",\"\",'{cl}'!A{cl_row})"
        for c in ("O", "P"):
            ws[f"{c}{r}"].font = Font(size=8, color=T["sample_text"])
    ws.column_dimensions["O"].hidden = True
    ws.column_dimensions["P"].hidden = True

    note = ws["B21"]
    note.value = ("Top-client list scans your first "
                  f"{eng_n} client rows. Hidden columns O-P do the math — "
                  "leave them alone (Format > Unhide if you're curious).")
    note.font = Font(italic=True, size=9, color=T["sample_text"])
    ws.merge_cells("B21:I21")


# ---------------------------------------------------------------------------
# START HERE
# ---------------------------------------------------------------------------

def build_start_here(wb):
    ws = wb[S["start"]]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = T["accent_dark"]
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 100
    ws.column_dimensions["C"].width = 3

    def head(row, text, big=False):
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = Font(bold=True, size=16 if big else 12, color=T["white"])
        cell.fill = fill(T["charcoal_dark"] if big else T["accent_dark"])
        cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 34 if big else 24

    def body(row, text, italic=False, height=None):
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = Font(size=11, italic=italic, color=T["charcoal_dark"])
        cell.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        cell.fill = fill(T["cream"])
        if height:
            ws.row_dimensions[row].height = height
        return cell

    r = 2
    head(r, f"WELCOME TO YOUR {CONFIG['product_name'].upper()}", big=True); r += 1
    body(r, "Made for dog groomers and grooming salons. Everything your "
            "business needs — clients, dogs, cut & style history, coat "
            "details, behavior warnings, appointments and money — in one "
            "place. No apps, no monthly software fee, nothing to install.",
         height=44); r += 2

    head(r, "STEP 1 — MAKE YOUR OWN COPY (do this first!)"); r += 1
    body(r, "USING GOOGLE SHEETS (free):  1. Go to sheets.google.com and sign "
            "in.  2. Click the folder icon (Open file picker), choose "
            "\"Upload\", and drop this file in.  3. It opens right in Google "
            "Sheets — then click File > Save as Google Sheets so it lives in "
            "your Drive forever.", height=58); r += 1
    body(r, "USING EXCEL:  just open the file and File > Save As with your "
            "business name. Done.", height=28); r += 2

    head(r, "STEP 2 — KNOW YOUR COLORS"); r += 1
    body(r, "WHITE cells = yours. Type anything.", height=20); r += 1
    c = body(r, "LIGHT GRAY cells = the spreadsheet's. They calculate "
                "themselves (rates, invoice totals, rebook due dates, the "
                "whole Dashboard). Don't type in them — if you ever do by "
                "accident, just press Ctrl+Z (Cmd+Z on Mac).", height=44)
    c.fill = fill(T["calc_fill"]); r += 1
    body(r, "Cells with a little arrow = dropdowns. Click the arrow and pick "
            "— never type these by hand, so names always match everywhere.",
         height=30); r += 2

    head(r, "STEP 3 — SET UP IN THIS ORDER"); r += 1
    body(r, "1. SERVICES & RATES — change the service names and prices to "
            "yours (full grooms by size, bath & brush, de-shed, nail trims). "
            "The Appointments tab fills in prices from this list "
            "automatically.", height=44); r += 1
    body(r, "2. CLIENTS — one row per owner. Phone, email, preferred contact, "
            "emergency contact, vet clinic and status all live here.",
         height=30); r += 1
    body(r, "3. PETS — one row per dog, and this tab is the whole point: "
            "coat type, last cut notes and exactly what the owner asked for, "
            "the blade and comb lengths you used, shampoo, skin conditions, "
            "matting spots, and ear / nail / gland preferences. Next visit, "
            "you open one row and you know the dog. Tip: type the word "
            "CAUTION anywhere in \"Behavior Notes\" and the whole row turns "
            "red — no one on your table gets surprised by a bite risk.",
         height=88); r += 1
    body(r, "4. APPOINTMENTS — one row per groom. Pick client and service "
            "from dropdowns; the rate appears by itself. Mark Status = Done "
            "when finished and Paid? = Yes when the money arrives. Today's "
            "appointments glow rose; finished-but-unpaid turn red. Type the "
            "rebook interval in weeks (most dogs: 4-8) and the \"Next Groom "
            "Due\" date calculates itself — it turns amber when it's time to "
            "call the client and rebook.", height=88); r += 1
    body(r, "5. INVOICES & PAYMENTS — type an invoice number, pick the "
            "client, and type the period start and end dates. The totals add "
            "themselves up from the Appointments tab (completed grooms in "
            "that period; Balance Due counts only the unpaid ones).",
         height=44); r += 2

    head(r, "THE SAMPLE ROWS"); r += 1
    body(r, "The gray italic rows (Amanda Foster, Derek Okafor, Priya "
            "Raman... and their dogs Murphy, Luna, Taco, Bella) are examples "
            "so you can see how everything works — including the red CAUTION "
            "row and the amber rebook reminder. When you're ready: click the "
            "row number on the left, drag to select the sample rows, "
            "right-click, and choose \"Delete rows\". The formulas and "
            "dropdowns stay — only the example text goes.", height=72); r += 2

    head(r, "ADDING MORE ROWS LATER"); r += 1
    body(r, "Formulas and dropdowns are already loaded far down each tab "
            "(Appointments to row 250, Invoices to row 100, Clients to row "
            "200, Pets to row 300). Just keep typing on the next empty line. "
            "If you ever fill a tab completely: select the last row, copy "
            "it, and paste it into the rows below — that carries the "
            "formulas down too.", height=58); r += 2

    head(r, "THE DASHBOARD"); r += 1
    body(r, "100% automatic. Active clients, dogs groomed this month, this "
            "week's appointments, today's appointments, unpaid balance, this "
            "month's revenue and your top clients — it all updates the "
            "moment you edit the other tabs.", height=44); r += 2

    body(r, "Questions? Message me on Etsy any time — happy to help.  Enjoy! ",
         italic=True); r += 1


# ============================================================================
# build
# ============================================================================

def build():
    wb = Workbook()
    order = ["start", "dash", "clients", "pets", "visits", "invoices",
             "services"]
    wb.active.title = S[order[0]]
    for key in order[1:]:
        wb.create_sheet(S[key])

    # defined names for cross-sheet dropdowns (oldest-Excel-safe + GSheets-safe)
    wb.defined_names.add(DefinedName(
        "ClientList",
        attr_text=f"'{S['clients']}'!$A$2:$A${R['clients_max']}"))
    wb.defined_names.add(DefinedName(
        "ServiceList",
        attr_text=f"'{S['services']}'!$A$2:$A${R['services_max']}"))

    build_start_here(wb)
    build_dashboard(wb)
    build_clients(wb)
    build_pets(wb)
    build_visits(wb)
    build_invoices(wb)
    build_services(wb)

    wb.calculation.fullCalcOnLoad = True
    out = os.path.join(HERE, CONFIG["output_file"])
    wb.save(out)
    return out


# ============================================================================
# verification
# ============================================================================

# Functions known to behave identically in Excel and Google Sheets.
ALLOWED_FUNCS = {
    "SUMIFS", "SUMIF", "COUNTIF", "COUNTIFS", "COUNTA", "COUNT", "SUM",
    "IF", "IFERROR", "AND", "OR", "NOT", "VLOOKUP", "INDEX", "MATCH",
    "LARGE", "SMALL", "TODAY", "DATE", "YEAR", "MONTH", "DAY", "WEEKDAY",
    "TEXT", "ROUND", "ROW", "COLUMN", "SEARCH", "ISNUMBER", "MIN", "MAX",
}

FUNC_RE = re.compile(r"([A-Z][A-Z0-9\.]*)\(")
SHEET_REF_RE = re.compile(r"'([^']+)'!")


def verify(path):
    problems = []
    wb = load_workbook(path)
    sheet_names = set(wb.sheetnames)
    n_formulas = 0

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                n_formulas += 1
                # strip string literals before analysis
                stripped = re.sub(r'"[^"]*"', '""', v)
                for fn in FUNC_RE.findall(stripped):
                    if fn not in ALLOWED_FUNCS:
                        problems.append(
                            f"{ws.title}!{cell.coordinate}: function {fn} "
                            f"not in dual-compat whitelist")
                for ref in SHEET_REF_RE.findall(stripped):
                    if ref not in sheet_names:
                        problems.append(
                            f"{ws.title}!{cell.coordinate}: reference to "
                            f"missing sheet '{ref}'")
                if stripped.count("(") != stripped.count(")"):
                    problems.append(
                        f"{ws.title}!{cell.coordinate}: unbalanced parens: {v}")

        for dv in ws.data_validations.dataValidation:
            f1 = dv.formula1 or ""
            if f1 in ("ClientList", "ServiceList"):
                continue
            if f1.startswith('"'):
                continue
            m = SHEET_REF_RE.match(f1)
            if m and m.group(1) not in sheet_names:
                problems.append(f"{ws.title} DV: missing sheet in {f1}")

    for name in ("ClientList", "ServiceList"):
        if name not in wb.defined_names:
            problems.append(f"missing defined name {name}")

    print(f"[verify/static] {n_formulas} formula cells scanned")
    return problems


def _lo_profile(tmp):
    """LibreOffice profile that forces recalculation of xlsx on load."""
    profile = os.path.join(tmp, "louser")
    os.makedirs(os.path.join(profile, "user"), exist_ok=True)
    with open(os.path.join(profile, "user",
                           "registrymodifications.xcu"), "w") as f:
        f.write(
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
            'xmlns:xs="http://www.w3.org/2001/XMLSchema">\n'
            '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
            '<prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
            '<value>0</value></prop></item>\n'
            '</oor:items>\n')
    return profile


def verify_recalc(path):
    """Recalculate with LibreOffice (recalc forced via profile setting) and
    look for error values in the computed results."""
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        profile = _lo_profile(tmp)
        cmd = ["soffice", "--headless", "--norestore",
               f"-env:UserInstallation=file://{profile}",
               "--convert-to", "xlsx", "--outdir", tmp, path]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        out = os.path.join(tmp, os.path.basename(path))
        if not os.path.exists(out):
            return [f"LibreOffice recalc failed: {res.stderr.strip()[:300]}"]
        wb = load_workbook(out, data_only=True)
        errs = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!",
                "#NULL!", "Err:")
        n_vals = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if v is None:
                        continue
                    n_vals += 1
                    if isinstance(v, str) and any(e in v for e in errs):
                        problems.append(
                            f"{ws.title}!{cell.coordinate}: computes to {v}")
        print(f"[verify/recalc] LibreOffice recalculated; "
              f"{n_vals} non-empty cells checked for error values")
    return problems


def spot_check(path):
    """Print a few computed values so a human can sanity-check the math."""
    with tempfile.TemporaryDirectory() as tmp:
        profile = _lo_profile(tmp)
        subprocess.run(["soffice", "--headless", "--norestore",
                        f"-env:UserInstallation=file://{profile}",
                        "--convert-to", "xlsx", "--outdir", tmp, path],
                       capture_output=True, text=True, timeout=300)
        wb = load_workbook(os.path.join(tmp, os.path.basename(path)),
                           data_only=True)
        d = wb[S["dash"]]
        print("[spot-check] Dashboard computed values:")
        for label, coord in [("Active Clients", "B6"),
                             ("Dogs Groomed This Month", "E6"),
                             ("Appointments This Week", "H6"),
                             ("Appointments Today", "B9"),
                             ("Revenue This Month", "E9"),
                             ("Unpaid Balance", "H9"),
                             ("Top client #1", "C14"),
                             ("Top client #1 revenue", "H14")]:
            print(f"    {label:24s} = {d[coord].value}")
        v = wb[S["visits"]]
        print("[spot-check] Appointment rates (VLOOKUP):",
              [v[f"E{r}"].value for r in range(2, 8)])
        print("[spot-check] Next-groom-due dates:",
              [v[f"I{r}"].value for r in range(2, 8)])
        inv = wb[S["invoices"]]
        print("[spot-check] Invoice totals/balances:",
              [(inv[f"E{r}"].value, inv[f"F{r}"].value) for r in range(2, 5)])


if __name__ == "__main__":
    out = build()
    print(f"built: {out}")
    probs = verify(out)
    probs += verify_recalc(out)
    if probs:
        print("PROBLEMS:")
        for p in probs:
            print("  -", p)
        sys.exit(1)
    spot_check(out)
    print("verification clean.")
