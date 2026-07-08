#!/usr/bin/env python3
"""
Cleaning Business Client CRM — Etsy digital product builder (re-skin of the
Pet Sitter / Dog Groomer CRM generator; see products/groomer_crm/build.py).

Generates Cleaning-Business-CRM.xlsx (Google Sheets + Excel compatible):
  * standard formulas only (SUMIFS, COUNTIF, COUNTIFS, IF, IFERROR, VLOOKUP,
    INDEX/MATCH, LARGE, TODAY, DATE, WEEKDAY, TEXT, ROUND, SEARCH, ISNUMBER)
    — no VBA, no macros, no Excel-only functions
  * data-validation dropdowns, formula-based conditional formatting
  * parameterized CONFIG so the same script can be re-skinned for other
    verticals by editing CONFIG only.

The moat for this niche is the trade-specific data free tools ignore:
property access (key location / lockbox / alarm / garage codes), pets on
site, who-supplies-the-products, per-home scope + recurring cadence, and a
"Next Clean Due" date that computes itself from that cadence.

Run:  python3 build.py          -> builds + verifies the workbook
"""

import os
import re
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))

# ============================================================================
# CONFIG — everything vertical-specific lives here.
# ============================================================================

CONFIG = {
    "output_file": "Cleaning-Business-CRM.xlsx",
    "product_name": "Cleaning Business Client CRM",
    "tagline": "Client, property-access, job, schedule & payment tracker "
               "for house cleaners",

    # ---- theme: steel blue + charcoal (distinct from sage petsitter / rose
    #      groomer) --------------------------------------------------------
    "theme": {
        "charcoal":      "2E3436",  # header rows (cool charcoal)
        "charcoal_dark": "232829",  # banners / title band
        "accent_dark":   "4A7196",  # steel blue — accents / section heads
        "accent_mid":    "6B93B3",  # tile labels / tab color
        "accent_light":  "E4EDF3",  # tile bodies / today-highlight (blue tint)
        "accent_pale":   "F1F6FA",  # page background feel (palest blue)
        "cream":         "FAFCFE",  # input cells (cool white)
        "calc_fill":     "ECEFF1",  # auto-calculated cells (do not type)
        "red_fill":      "F6D9D5",
        "red_text":      "9C2B23",
        "amber_fill":    "FBEFD0",
        "amber_text":    "8A6116",
        "green_fill":    "DDEBD8",
        "green_text":    "3E6B47",
        "sample_text":   "6D6D6D",  # sample rows: gray italic
        "white":         "FFFFFF",
    },

    # ---- sheet display names (order = tab order) -------------------------
    "sheet_names": {
        "start":    "START HERE",
        "dash":     "Dashboard",
        "clients":  "Clients",
        "jobs":     "Job Details",
        "schedule": "Schedule",
        "invoices": "Invoices & Payments",
        "services": "Services & Rates",
    },

    # ---- vertical vocabulary ----------------------------------------------
    "vocab": {
        "client": "Client",
        "unit": "Home",
        "unit_plural": "Homes",
        "visit": "Clean",
        "caution_keyword": "CAUTION",
    },

    # ---- dropdown option lists (inline lists, comma-safe values) ---------
    "lists": {
        "client_status": ["Active", "Inactive", "Waitlist"],
        "contact_method": ["Text", "Call", "Email"],
        "supplies": ["Client supplies", "I supply", "Mixed"],
        "cadence": ["Weekly", "Biweekly", "Monthly", "One-off", "Deep Clean"],
        "visit_status": ["Booked", "Done", "Cancelled"],
        "paid": ["Yes", "No"],
        "invoice_status": ["Draft", "Sent", "Paid", "Overdue"],
    },

    # ---- editable price list (Services & Rates tab) -----------------------
    # These are NOT sample rows — they're the real starter price list.
    "services": [
        ("Standard Clean",          120.00, "Flat / ~3 hrs",
         "Kitchen, baths, floors, dusting, trash — a maintained home"),
        ("Deep Clean",              250.00, "Flat / 5-6 hrs",
         "Baseboards, inside windows, cabinet fronts, scrub build-up"),
        ("Move-Out / Move-In",      320.00, "Flat / 6-8 hrs",
         "Empty home: inside cabinets, oven, fridge, closets, windows"),
        ("Recurring Clean",         100.00, "Per visit",
         "Standard clean for weekly / biweekly / monthly regulars"),
        ("Hourly",                   45.00, "Per hour",
         "Odd jobs or unknown scope — bill the actual time worked"),
        ("Add-on: Inside Fridge",    35.00, "Per item",
         "Empty, wipe and sanitize the interior"),
        ("Add-on: Inside Oven",      35.00, "Per item",
         "Degrease racks and interior"),
        ("Add-on: Interior Windows", 45.00, "Per item",
         "Interior glass + sills, standard home"),
        ("Add-on: Inside Cabinets",  40.00, "Per item",
         "Empty kitchen cabinets wiped inside"),
        ("Add-on: Laundry / Linens", 25.00, "Per load",
         "Wash, dry and fold, or strip & remake beds"),
    ],

    # ---- column definitions ------------------------------------------------
    # type: input | dropdown:<listkey> | dropdown_range:<defined name> |
    #       date | date_calc | currency_calc | currency_input
    "columns": {
        "clients": [
            ("Client Name",               22, "input"),
            ("Phone",                     15, "input"),
            ("Email",                     24, "input"),
            ("Service Address",           30, "input"),
            ("Entry / Key Location",      27, "input"),
            ("Lockbox Code",              13, "input"),
            ("Alarm Code",                15, "input"),
            ("Garage Code",               12, "input"),
            ("Pets on Site",              26, "input"),
            ("Products / Allergies",      28, "input"),
            ("Who Supplies",              15, "dropdown:supplies"),
            ("Preferred Contact",         16, "dropdown:contact_method"),
            ("Status",                    12, "dropdown:client_status"),
            ("Notes",                     30, "input"),
        ],
        "jobs": [
            ("Client",                    22, "dropdown_range:ClientList"),
            ("Bedrooms",                  10, "input"),
            ("Bathrooms",                 10, "input"),
            ("Sq Ft",                      9, "input"),
            ("Scope / Rooms Checklist",   42, "input"),
            ("Recurring Cadence",         16, "dropdown:cadence"),
            ("Flat Rate ($)",             12, "currency_input"),
            ("Hourly Rate ($)",           13, "currency_input"),
            ("Est. Hours",                10, "input"),
            ("Notes",                     30, "input"),
        ],
        "schedule": [
            ("Date",                      12, "date"),
            ("Client",                    22, "dropdown_range:ClientList"),
            ("Service",                   22, "dropdown_range:ServiceList"),
            ("Rate ($)",                  10, "currency_calc"),
            ("Status",                    11, "dropdown:visit_status"),
            ("Paid?",                      8, "dropdown:paid"),
            ("Next Clean Due",            15, "date_calc"),
            ("Notes",                     30, "input"),
        ],
        "invoices": [
            ("Invoice #",                 11, "input"),
            ("Client",                    22, "dropdown_range:ClientList"),
            ("Period Start",              13, "date"),
            ("Period End",                13, "date"),
            ("Invoice Total ($)",         15, "currency_calc"),
            ("Balance Due ($)",           15, "currency_calc"),
            ("Date Sent",                 12, "date"),
            ("Date Paid",                 12, "date"),
            ("Status",                    11, "dropdown:invoice_status"),
            ("Notes",                     30, "input"),
        ],
        "services": [
            ("Service",                   24, "input"),
            ("Rate ($)",                  11, "currency_input"),
            ("Unit / Time",               14, "input"),
            ("Notes",                     46, "input"),
        ],
    },

    # ---- how deep formulas / dropdowns are pre-loaded ----------------------
    "rows": {
        "clients_max": 200,    # ClientList range + dropdown depth
        "services_max": 30,    # ServiceList range
        "jobs_max": 250,       # Job Details depth
        "schedule_max": 250,   # rate + next-due formulas pre-filled to here
        "invoices_max": 100,   # SUMIFS pre-filled to here
        "top_clients": 5,      # dashboard top-N
        "top_engine": 20,      # first N client rows scanned for top clients
    },

    # ---- sample rows (gray italic, "delete me" note, listed in START HERE) --
    # Dates are TODAY()-relative formulas so the demo conditional formatting
    # (today's cleans, done+unpaid, next-clean-due) actually lights up.
    "sample": {
        "clients": [
            ["Sofia Delgado", "(555) 233-9087", "sofia.d@example.com",
             "418 Maple Ct, Riverside",
             "Side gate — key under blue planter on back porch",
             "N/A", "2255 #", "1490",
             "Golden retriever 'Biscuit' — friendly, keep in the yard",
             "Client supplies; fragrance-free only (asthma)", "Client supplies",
             "Text", "Active", "Weekly, Fri mornings — SAMPLE ROW, delete me"],
            ["Marcus Reilly", "(555) 661-2048", "marcus.r@example.com",
             "77 Birchwood Ln, Apt 3B",
             "Front desk holds a spare key — ask for unit 3B",
             "8830", "none", "none",
             "No pets", "I bring all supplies", "I supply",
             "Call", "Active",
             "Biweekly — elevator building — SAMPLE ROW, delete me"],
            ["Priya Nadar", "(555) 902-7715", "priya.n@example.com",
             "1205 Kestrel Dr",
             "Lockbox on the hose bib, back patio",
             "5127", "4400 (code word: Willow)", "2213",
             "CAUTION: guard dog 'Rex' — owner MUST crate before each visit; "
             "do not enter the yard until confirmed",
             "Mixed — I bring supplies, client provides hardwood cleaner",
             "Mixed", "Email", "Active",
             "Monthly deep clean — SAMPLE ROW, delete me"],
            ["Aaron Brooks", "(555) 448-1190", "aaron.b@example.com",
             "9 Sutter Pl",
             "Meet on site for the first walkthrough",
             "TBD", "TBD", "TBD",
             "Dog — confirm at walkthrough", "TBD at walkthrough",
             "Client supplies", "Text", "Waitlist",
             "Move-out quote requested — SAMPLE ROW, delete me"],
        ],
        "jobs": [
            ["Sofia Delgado", "4", "3", "2400",
             "Whole house; hardwood + tile throughout; skip the garage",
             "Weekly", 140.00, 45.00, "3",
             "Fragrance-free products only — SAMPLE ROW, delete me"],
            ["Marcus Reilly", "2", "2", "1100",
             "Kitchen, 2 baths, living, 2 bedrooms; fold in-unit laundry",
             "Biweekly", 110.00, 45.00, "2.5",
             "Unit 3B, elevator building — SAMPLE ROW, delete me"],
            ["Priya Nadar", "3", "2", "1850",
             "Deep rotation: baseboards + inside windows monthly; "
             "confirm dog is crated first",
             "Monthly", 220.00, 50.00, "4",
             "Guard dog — confirm crated on arrival — SAMPLE ROW, delete me"],
            ["Aaron Brooks", "3", "2", "1600",
             "Move-out: inside cabinets, oven, fridge, interior windows",
             "One-off", 300.00, 50.00, "6",
             "Quote pending walkthrough — SAMPLE ROW, delete me"],
        ],
        # date offset (days from TODAY, or a literal formula string),
        # client, service, status, paid, note
        "schedule": [
            [0,   "Sofia Delgado", "Recurring Clean", "Booked", "No",
             "SAMPLE — today's cleans glow blue"],
            [-3,  "Marcus Reilly", "Recurring Clean", "Done",   "No",
             "SAMPLE — done + unpaid turns red"],
            [-10, "Sofia Delgado", "Recurring Clean", "Done",   "Yes",
             "SAMPLE — Next Clean Due turns amber once it is due"],
            # anchored to the 1st of the current month so the "Revenue This
            # Month" dashboard tile always has something to show
            ["=DATE(YEAR(TODAY()),MONTH(TODAY()),1)",
             "Sofia Delgado", "Deep Clean", "Done", "Yes",
             "SAMPLE ROW — delete me"],
            [-35, "Priya Nadar", "Deep Clean", "Done", "No",
             "SAMPLE — older unpaid clean (see Invoices) — delete me"],
            [2,   "Marcus Reilly", "Standard Clean", "Booked", "No",
             "SAMPLE ROW — delete me"],
        ],
        # inv#, client, period_start_offset, period_end_offset,
        # sent_offset (None = blank), paid_offset (None = blank), status, note
        "invoices": [
            ["INV-2001", "Sofia Delgado", -40, 1, -30, -25, "Paid",
             "SAMPLE ROW — delete me"],
            ["INV-2002", "Marcus Reilly", -8, 3, 0, None, "Sent",
             "SAMPLE ROW — delete me"],
            ["INV-2003", "Priya Nadar", -45, -2, -18, None, "Overdue",
             "SAMPLE ROW — delete me"],
        ],
    },
}

T = CONFIG["theme"]
S = CONFIG["sheet_names"]
R = CONFIG["rows"]

# ============================================================================
# style helpers
# ============================================================================

def fill(hex_):
    return PatternFill("solid", fgColor=hex_)

THIN = Side(style="thin", color="D8E1E8")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_FONT = Font(name="Calibri", bold=True, color=T["white"], size=11)
HEADER_FILL = fill(T["charcoal"])
SAMPLE_FONT = Font(name="Calibri", italic=True, color=T["sample_text"], size=11)
BODY_FONT = Font(name="Calibri", size=11)

CURRENCY_FMT = '"$"#,##0.00'
DATE_FMT = "mm/dd/yyyy"


def style_header_row(ws, n_cols, row=1, height=26):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[row].height = height


def set_widths(ws, cols):
    for i, (_, width, _) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def apply_col_formats(ws, cols, first_row, last_row):
    """Number formats + calc-cell shading for whole data region."""
    for i, (_, _, ctype) in enumerate(cols, start=1):
        for r in range(first_row, last_row + 1):
            cell = ws.cell(row=r, column=i)
            if ctype in ("currency_calc", "currency_input"):
                cell.number_format = CURRENCY_FMT
            elif ctype in ("date", "date_calc"):
                cell.number_format = DATE_FMT
            if ctype in ("currency_calc", "date_calc"):
                cell.fill = fill(T["calc_fill"])


def mark_sample_rows(ws, n_cols, first_row, n_rows):
    for r in range(first_row, first_row + n_rows):
        for c in range(1, n_cols + 1):
            ws.cell(row=r, column=c).font = SAMPLE_FONT


def add_dv(ws, dv_type, formula1, cell_range):
    dv = DataValidation(type=dv_type, formula1=formula1, allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv


def inline_list(key):
    return '"' + ",".join(CONFIG["lists"][key]) + '"'


def col_letter(sheet_key, label):
    """Column letter for a labeled column on a data sheet."""
    labels = [l for l, _, _ in CONFIG["columns"][sheet_key]]
    return get_column_letter(labels.index(label) + 1)


def build_columns(ws, sheet_key, max_row):
    """Headers, widths, dropdown DV for a standard data sheet."""
    cols = CONFIG["columns"][sheet_key]
    for i, (label, _, _) in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=label)
    style_header_row(ws, len(cols))
    set_widths(ws, cols)
    ws.freeze_panes = "A2"
    for i, (_, _, ctype) in enumerate(cols, start=1):
        col = get_column_letter(i)
        rng = f"{col}2:{col}{max_row}"
        if ctype.startswith("dropdown_range:"):
            add_dv(ws, "list", ctype.split(":", 1)[1], rng)
        elif ctype.startswith("dropdown:"):
            add_dv(ws, "list", inline_list(ctype.split(":", 1)[1]), rng)
    apply_col_formats(ws, cols, 2, max_row)
    return cols


# ============================================================================
# sheet builders
# ============================================================================

def build_services(wb):
    ws = wb[S["services"]]
    cols = build_columns(ws, "services", R["services_max"])
    for r, row in enumerate(CONFIG["services"], start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    # gentle banding on the real price rows
    for r in range(2, len(CONFIG["services"]) + 2):
        for c in range(1, len(cols) + 1):
            ws.cell(row=r, column=c).border = BORDER_ALL
    note = ws.cell(row=R["services_max"] + 2, column=1,
                   value="Edit names and rates freely — the Schedule tab looks "
                         "prices up from this list automatically. Service names "
                         "here must match the dropdown picks exactly (they "
                         "will, if you always pick from the dropdown).")
    note.font = Font(italic=True, color=T["sample_text"], size=10)
    ws.merge_cells(start_row=R["services_max"] + 2, start_column=1,
                   end_row=R["services_max"] + 2, end_column=4)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_properties.tabColor = T["accent_dark"]


def build_clients(wb):
    ws = wb[S["clients"]]
    cols = build_columns(ws, "clients", R["clients_max"])
    samples = CONFIG["sample"]["clients"]
    for r, row in enumerate(samples, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    # status column CF: Active = green tint, Waitlist = amber (added FIRST so
    # the status cell keeps its color even inside a CAUTION-red row)
    status_col = col_letter("clients", "Status")
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{R['clients_max']}",
        FormulaRule(formula=[f'${status_col}2="Active"'],
                    fill=fill(T["green_fill"]),
                    font=Font(color=T["green_text"], bold=True)))
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{R['clients_max']}",
        FormulaRule(formula=[f'${status_col}2="Waitlist"'],
                    fill=fill(T["amber_fill"]),
                    font=Font(color=T["amber_text"], bold=True)))
    # CAUTION keyword in Pets-on-Site OR Notes -> whole row red (safety alert:
    # loose dogs, finicky alarms, hazards you must not miss on arrival)
    pets_col = col_letter("clients", "Pets on Site")
    notes_col = col_letter("clients", "Notes")
    last_col = get_column_letter(len(cols))
    kw = CONFIG["vocab"]["caution_keyword"]
    ws.conditional_formatting.add(
        f"A2:{last_col}{R['clients_max']}",
        FormulaRule(
            formula=[f'OR(ISNUMBER(SEARCH("{kw}",${pets_col}2)),'
                     f'ISNUMBER(SEARCH("{kw}",${notes_col}2)))'],
            fill=fill(T["red_fill"]),
            font=Font(color=T["red_text"]),
            stopIfTrue=True))
    ws.sheet_properties.tabColor = T["accent_mid"]


def build_jobs(wb):
    ws = wb[S["jobs"]]
    cols = build_columns(ws, "jobs", R["jobs_max"])
    samples = CONFIG["sample"]["jobs"]
    for r, row in enumerate(samples, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    apply_col_formats(ws, cols, 2, R["jobs_max"])
    ws.sheet_properties.tabColor = T["accent_mid"]


def build_schedule(wb):
    ws = wb[S["schedule"]]
    cols = build_columns(ws, "schedule", R["schedule_max"])
    labels = [l for l, _, _ in cols]
    c_date = col_letter("schedule", "Date")            # A
    c_client = col_letter("schedule", "Client")        # B
    c_service = col_letter("schedule", "Service")      # C
    c_status = col_letter("schedule", "Status")        # E
    c_paid = col_letter("schedule", "Paid?")           # F
    c_due = col_letter("schedule", "Next Clean Due")   # G
    last_col = get_column_letter(len(cols))
    svc_sheet = S["services"]
    jobs_sheet = S["jobs"]
    jobs_last = get_column_letter(len(CONFIG["columns"]["jobs"]))
    cad_idx = [l for l, _, _ in CONFIG["columns"]["jobs"]].index(
        "Recurring Cadence") + 1
    jobs_rng = (f"'{jobs_sheet}'!$A$2:${jobs_last}${R['jobs_max']}")

    # pre-load the rate lookup (from Services) + the "Next Clean Due" date,
    # which reads the home's recurring cadence off the Job Details tab and
    # adds the matching number of days to this clean's date.
    for r in range(2, R["schedule_max"] + 1):
        ws.cell(row=r, column=labels.index("Rate ($)") + 1).value = (
            f'=IF(${c_service}{r}="","",'
            f"IFERROR(VLOOKUP(${c_service}{r},"
            f"'{svc_sheet}'!$A$2:$B${R['services_max']},2,FALSE),\"\"))"
        )
        cad = f"VLOOKUP(${c_client}{r},{jobs_rng},{cad_idx},FALSE)"
        ws.cell(row=r, column=labels.index("Next Clean Due") + 1).value = (
            f'=IF(OR(${c_date}{r}="",${c_client}{r}=""),"",'
            f'IFERROR(IF({cad}="Weekly",${c_date}{r}+7,'
            f'IF({cad}="Biweekly",${c_date}{r}+14,'
            f'IF({cad}="Monthly",${c_date}{r}+30,""))),""))'
        )

    samples = CONFIG["sample"]["schedule"]
    for r, (off, client, service, status, paid, note) in \
            enumerate(samples, start=2):
        if isinstance(off, str):
            date_formula = off
        else:
            date_formula = (f"=TODAY()+{off}" if off >= 0
                            else f"=TODAY()-{abs(off)}")
        ws.cell(row=r, column=labels.index("Date") + 1,
                value=date_formula).number_format = DATE_FMT
        ws.cell(row=r, column=labels.index("Client") + 1, value=client)
        ws.cell(row=r, column=labels.index("Service") + 1, value=service)
        ws.cell(row=r, column=labels.index("Status") + 1, value=status)
        ws.cell(row=r, column=labels.index("Paid?") + 1, value=paid)
        ws.cell(row=r, column=labels.index("Notes") + 1, value=note)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    # keep number formats + calc shading after sample write
    apply_col_formats(ws, cols, 2, R["schedule_max"])

    # CF 1 (wins): done + not paid -> red row
    ws.conditional_formatting.add(
        f"A2:{last_col}{R['schedule_max']}",
        FormulaRule(
            formula=[f'AND(${c_status}2="Done",${c_paid}2<>"Yes")'],
            fill=fill(T["red_fill"]), font=Font(color=T["red_text"]),
            stopIfTrue=True))
    # CF 2: today's cleans -> steel-blue highlight
    ws.conditional_formatting.add(
        f"A2:{last_col}{R['schedule_max']}",
        FormulaRule(formula=[f"${c_date}2=TODAY()"],
                    fill=fill(T["accent_light"]),
                    font=Font(color=T["charcoal_dark"], bold=True),
                    stopIfTrue=True))
    # CF 3: next-clean-due date has arrived -> amber on the Next Clean Due cell
    ws.conditional_formatting.add(
        f"{c_due}2:{c_due}{R['schedule_max']}",
        FormulaRule(
            formula=[f'AND(ISNUMBER(${c_due}2),${c_due}2<=TODAY())'],
            fill=fill(T["amber_fill"]),
            font=Font(color=T["amber_text"], bold=True)))
    ws.sheet_properties.tabColor = T["accent_mid"]


def build_invoices(wb):
    ws = wb[S["invoices"]]
    cols = build_columns(ws, "invoices", R["invoices_max"])
    labels = [l for l, _, _ in cols]
    c_client = col_letter("invoices", "Client")
    c_start = col_letter("invoices", "Period Start")
    c_end = col_letter("invoices", "Period End")
    c_status = col_letter("invoices", "Status")
    last_col = get_column_letter(len(cols))
    sc = S["schedule"]
    s_rate = col_letter("schedule", "Rate ($)")
    s_client = col_letter("schedule", "Client")
    s_date = col_letter("schedule", "Date")
    s_status = col_letter("schedule", "Status")
    s_paid = col_letter("schedule", "Paid?")

    for r in range(2, R["invoices_max"] + 1):
        common = (f"'{sc}'!${s_rate}:${s_rate},"
                  f"'{sc}'!${s_client}:${s_client},${c_client}{r},"
                  f"'{sc}'!${s_status}:${s_status},\"Done\","
                  f"'{sc}'!${s_date}:${s_date},\">=\"&${c_start}{r},"
                  f"'{sc}'!${s_date}:${s_date},\"<=\"&${c_end}{r}")
        guard = (f'OR(${c_client}{r}="",${c_start}{r}="",${c_end}{r}="")')
        ws.cell(row=r, column=labels.index("Invoice Total ($)") + 1).value = \
            f'=IF({guard},"",SUMIFS({common}))'
        ws.cell(row=r, column=labels.index("Balance Due ($)") + 1).value = \
            (f'=IF({guard},"",SUMIFS({common},'
             f"'{sc}'!${s_paid}:${s_paid},\"<>Yes\"))")

    def off_formula(off):
        if off is None:
            return None
        return f"=TODAY()+{off}" if off >= 0 else f"=TODAY()-{abs(off)}"

    samples = CONFIG["sample"]["invoices"]
    for r, (num, client, ps, pe, sent, paid, status, note) in \
            enumerate(samples, start=2):
        ws.cell(row=r, column=1, value=num)
        ws.cell(row=r, column=2, value=client)
        ws.cell(row=r, column=3, value=off_formula(ps))
        ws.cell(row=r, column=4, value=off_formula(pe))
        if sent is not None:
            ws.cell(row=r, column=7, value=off_formula(sent))
        if paid is not None:
            ws.cell(row=r, column=8, value=off_formula(paid))
        ws.cell(row=r, column=9, value=status)
        ws.cell(row=r, column=10, value=note)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    apply_col_formats(ws, cols, 2, R["invoices_max"])

    # status CF on whole row
    rng = f"A2:{last_col}{R['invoices_max']}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'${c_status}2="Overdue"'], fill=fill(T["red_fill"]),
        font=Font(color=T["red_text"]), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'${c_status}2="Paid"'], fill=fill(T["green_fill"]),
        font=Font(color=T["green_text"]), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'${c_status}2="Sent"'], fill=fill(T["amber_fill"]),
        font=Font(color=T["amber_text"]), stopIfTrue=True))
    ws.sheet_properties.tabColor = T["accent_mid"]


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

def _tile(ws, col, label, formula, num_fmt, label_row, value_row):
    """Two merged 2-wide cells stacked: label on top, big value below."""
    c1, c2 = col, get_column_letter(column_index_from_string(col) + 1)
    ws.merge_cells(f"{c1}{label_row}:{c2}{label_row}")
    ws.merge_cells(f"{c1}{value_row}:{c2}{value_row}")
    lab = ws[f"{c1}{label_row}"]
    lab.value = label.upper()
    lab.font = Font(bold=True, size=9, color=T["white"])
    lab.fill = fill(T["accent_dark"])
    lab.alignment = Alignment(horizontal="center", vertical="center")
    val = ws[f"{c1}{value_row}"]
    val.value = formula
    val.font = Font(bold=True, size=22, color=T["charcoal_dark"])
    val.fill = fill(T["accent_light"])
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.number_format = num_fmt
    for cc in (c1, c2):
        ws[f"{cc}{label_row}"].fill = fill(T["accent_dark"])
        ws[f"{cc}{value_row}"].fill = fill(T["accent_light"])


def build_dashboard(wb):
    ws = wb[S["dash"]]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = T["charcoal_dark"]

    cl, sc = S["clients"], S["schedule"]
    cl_status = col_letter("clients", "Status")
    s_date = col_letter("schedule", "Date")
    s_client = col_letter("schedule", "Client")
    s_rate = col_letter("schedule", "Rate ($)")
    s_status = col_letter("schedule", "Status")
    s_paid = col_letter("schedule", "Paid?")

    # widths: A spacer, B..I tiles, then helper cols
    ws.column_dimensions["A"].width = 2
    for c in "BCEFHI":
        ws.column_dimensions[c].width = 13
    for c in "DG":
        ws.column_dimensions[c].width = 2
    ws.column_dimensions["J"].width = 2
    ws.column_dimensions["O"].width = 12
    ws.column_dimensions["P"].width = 18

    # title band
    ws.merge_cells("B2:I2")
    t = ws["B2"]
    t.value = CONFIG["product_name"].upper() + "  —  DASHBOARD"
    t.font = Font(bold=True, size=16, color=T["white"])
    t.fill = fill(T["charcoal_dark"])
    t.alignment = Alignment(horizontal="center", vertical="center")
    for c in "CDEFGHI":
        ws[f"{c}2"].fill = fill(T["charcoal_dark"])
    ws.row_dimensions[2].height = 30

    ws["B3"] = ('="Everything on this page updates itself — nothing to type '
                'here.  Today: "&TEXT(TODAY(),"mmm d, yyyy")')
    ws["B3"].font = Font(italic=True, size=10, color=T["sample_text"])

    week_start = "(TODAY()-WEEKDAY(TODAY(),2)+1)"
    month_start = "DATE(YEAR(TODAY()),MONTH(TODAY()),1)"
    next_month = "DATE(YEAR(TODAY()),MONTH(TODAY())+1,1)"
    tiles = [
        ("B", "Active Clients",
         f"=COUNTIF('{cl}'!${cl_status}:${cl_status},\"Active\")", "0"),
        ("E", "Cleans This Month",
         f"=COUNTIFS('{sc}'!${s_status}:${s_status},\"Done\","
         f"'{sc}'!${s_date}:${s_date},\">=\"&{month_start},"
         f"'{sc}'!${s_date}:${s_date},\"<\"&{next_month})", "0"),
        ("H", "Cleans This Week",
         f"=COUNTIFS('{sc}'!${s_date}:${s_date},\">=\"&{week_start},"
         f"'{sc}'!${s_date}:${s_date},\"<=\"&{week_start}+6,"
         f"'{sc}'!${s_status}:${s_status},\"<>Cancelled\")", "0"),
    ]
    tiles2 = [
        ("B", "Cleans Today",
         f"=COUNTIFS('{sc}'!${s_date}:${s_date},TODAY(),"
         f"'{sc}'!${s_status}:${s_status},\"<>Cancelled\")", "0"),
        ("E", "Revenue This Month",
         f"=SUMIFS('{sc}'!${s_rate}:${s_rate},"
         f"'{sc}'!${s_status}:${s_status},\"Done\","
         f"'{sc}'!${s_date}:${s_date},\">=\"&{month_start},"
         f"'{sc}'!${s_date}:${s_date},\"<\"&{next_month})",
         CURRENCY_FMT),
        ("H", "Unpaid Balance",
         f"=SUMIFS('{sc}'!${s_rate}:${s_rate},"
         f"'{sc}'!${s_status}:${s_status},\"Done\","
         f"'{sc}'!${s_paid}:${s_paid},\"<>Yes\")",
         CURRENCY_FMT),
    ]
    for col, label, formula, fmt in tiles:
        _tile(ws, col, label, formula, fmt, 5, 6)
    for col, label, formula, fmt in tiles2:
        _tile(ws, col, label, formula, fmt, 8, 9)
    ws.row_dimensions[6].height = 40
    ws.row_dimensions[9].height = 40

    # ---- top clients table -------------------------------------------------
    ws.merge_cells("B12:I12")
    h = ws["B12"]
    h.value = "TOP CLIENTS  (completed cleans, all time)"
    h.font = Font(bold=True, size=12, color=T["white"])
    h.fill = fill(T["charcoal"])
    h.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in "CDEFGHI":
        ws[f"{c}12"].fill = fill(T["charcoal"])
    ws.row_dimensions[12].height = 22

    hdr = {"B": "Rank", "C": "Client", "H": "Revenue"}
    ws.merge_cells("C13:G13")
    for c, v in hdr.items():
        cell = ws[f"{c}13"]
        cell.value = v
        cell.font = Font(bold=True, size=10, color=T["accent_dark"])
        cell.fill = fill(T["accent_pale"])
    ws.merge_cells("H13:I13")
    ws["H13"].fill = fill(T["accent_pale"])

    eng_top, eng_n = 14, R["top_engine"]           # engine rows 14..33
    eng_rng = f"$O${eng_top}:$O${eng_top + eng_n - 1}"
    name_rng = f"$P${eng_top}:$P${eng_top + eng_n - 1}"
    for k in range(1, R["top_clients"] + 1):
        row = 13 + k
        ws[f"B{row}"] = k
        ws[f"B{row}"].alignment = Alignment(horizontal="center")
        ws.merge_cells(f"C{row}:G{row}")
        ws[f"C{row}"] = (f"=IFERROR(INDEX({name_rng},"
                         f"MATCH(LARGE({eng_rng},{k}),{eng_rng},0)),\"—\")")
        ws.merge_cells(f"H{row}:I{row}")
        ws[f"H{row}"] = f'=IFERROR(ROUND(LARGE({eng_rng},{k}),2),"")'
        ws[f"H{row}"].number_format = CURRENCY_FMT
        for c in "BCDEFGHI":
            ws[f"{c}{row}"].fill = fill(
                T["accent_pale"] if k % 2 else T["white"])
        ws[f"C{row}"].font = Font(size=11, color=T["charcoal_dark"])
        ws[f"H{row}"].font = Font(size=11, bold=True, color=T["accent_dark"])

    # ---- hidden-ish calculation engine (cols O:P) ---------------------------
    ws["O13"] = "engine ↓"
    ws["P13"] = "do not edit"
    for c in ("O13", "P13"):
        ws[c].font = Font(size=8, color=T["sample_text"], italic=True)
    for i in range(eng_n):
        r = eng_top + i
        cl_row = 2 + i
        ws[f"O{r}"] = (f"=IF('{cl}'!A{cl_row}=\"\",\"\","
                       f"SUMIFS('{sc}'!${s_rate}:${s_rate},"
                       f"'{sc}'!${s_client}:${s_client},'{cl}'!A{cl_row},"
                       f"'{sc}'!${s_status}:${s_status},\"Done\")"
                       f"+ROW()/10000000)")
        ws[f"P{r}"] = f"=IF($O{r}=\"\",\"\",'{cl}'!A{cl_row})"
        for c in ("O", "P"):
            ws[f"{c}{r}"].font = Font(size=8, color=T["sample_text"])
    ws.column_dimensions["O"].hidden = True
    ws.column_dimensions["P"].hidden = True

    note = ws["B21"]
    note.value = ("Top-client list scans your first "
                  f"{eng_n} client rows. Hidden columns O-P do the math — "
                  "leave them alone (Format > Unhide if you're curious).")
    note.font = Font(italic=True, size=9, color=T["sample_text"])
    ws.merge_cells("B21:I21")


# ---------------------------------------------------------------------------
# START HERE
# ---------------------------------------------------------------------------

def build_start_here(wb):
    ws = wb[S["start"]]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = T["accent_dark"]
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 100
    ws.column_dimensions["C"].width = 3

    def head(row, text, big=False):
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = Font(bold=True, size=16 if big else 12, color=T["white"])
        cell.fill = fill(T["charcoal_dark"] if big else T["accent_dark"])
        cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 34 if big else 24

    def body(row, text, italic=False, height=None):
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = Font(size=11, italic=italic, color=T["charcoal_dark"])
        cell.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        cell.fill = fill(T["cream"])
        if height:
            ws.row_dimensions[row].height = height
        return cell

    r = 2
    head(r, f"WELCOME TO YOUR {CONFIG['product_name'].upper()}", big=True); r += 1
    body(r, "Made for house cleaners, maid services and cleaning companies. "
            "Everything your business needs — clients, how to get into each "
            "home, the scope of each job, your schedule and your money — in "
            "one place. No apps, no monthly software fee, nothing to install.",
         height=44); r += 2

    head(r, "STEP 1 — MAKE YOUR OWN COPY (do this first!)"); r += 1
    body(r, "USING GOOGLE SHEETS (free):  1. Go to sheets.google.com and sign "
            "in.  2. Click the folder icon (Open file picker), choose "
            "\"Upload\", and drop this file in.  3. It opens right in Google "
            "Sheets — then click File > Save as Google Sheets so it lives in "
            "your Drive forever.", height=58); r += 1
    body(r, "USING EXCEL:  just open the file and File > Save As with your "
            "business name. Done.", height=28); r += 2

    head(r, "STEP 2 — KNOW YOUR COLORS"); r += 1
    body(r, "WHITE cells = yours. Type anything.", height=20); r += 1
    c = body(r, "LIGHT GRAY cells = the spreadsheet's. They calculate "
                "themselves (the rate on each clean, invoice totals, the Next "
                "Clean Due date, the whole Dashboard). Don't type in them — if "
                "you ever do by accident, just press Ctrl+Z (Cmd+Z on Mac).",
             height=44)
    c.fill = fill(T["calc_fill"]); r += 1
    body(r, "Cells with a little arrow = dropdowns. Click the arrow and pick "
            "— never type these by hand, so names always match everywhere.",
         height=30); r += 2

    head(r, "STEP 3 — SET UP IN THIS ORDER"); r += 1
    body(r, "1. SERVICES & RATES — change the service names and prices to "
            "yours (Standard Clean, Deep Clean, Move-Out, your Recurring rate, "
            "your hourly rate, and add-ons like inside fridge / oven / windows). "
            "The Schedule tab fills in the price from this list automatically.",
         height=44); r += 1
    body(r, "2. CLIENTS — one row per client. Phone, email, the service "
            "address, and the part free tools never have room for: HOW YOU GET "
            "IN. Key location, lockbox code, alarm code and garage code each "
            "get their own column. Also: pets on site, who supplies the "
            "products (and any allergies / scent rules), preferred contact and "
            "an Active / Inactive / Waitlist status. Tip: type the word CAUTION "
            "anywhere in \"Pets on Site\" or \"Notes\" and the whole row turns "
            "red — a loose dog or a finicky alarm is the last thing you want to "
            "forget on the doorstep.", height=100); r += 1
    body(r, "3. JOB DETAILS — one row per home: bedrooms, bathrooms, square "
            "footage, your room-by-room scope, and the RECURRING CADENCE "
            "dropdown (Weekly / Biweekly / Monthly / One-off / Deep Clean). "
            "Set the flat rate and/or your hourly rate for that home. The "
            "cadence you pick here is what drives the \"Next Clean Due\" date "
            "over on the Schedule.", height=72); r += 1
    body(r, "4. SCHEDULE — one row per clean. Pick the client and the service "
            "from dropdowns; the rate appears by itself. Mark Status = Done "
            "when the job is finished and Paid? = Yes when the money arrives. "
            "Today's cleans glow blue; finished-but-unpaid cleans turn red. The "
            "\"Next Clean Due\" date fills itself in from that home's cadence "
            "and turns amber when it's time to book the next visit.",
         height=86); r += 1
    body(r, "5. INVOICES & PAYMENTS — type an invoice number, pick the client, "
            "and type the period start and end dates. The totals add "
            "themselves up from the Schedule tab (completed cleans in that "
            "period; Balance Due counts only the unpaid ones).",
         height=44); r += 2

    head(r, "A NOTE ON THE ACCESS CODES"); r += 1
    body(r, "This file holds lockbox, alarm and garage codes, so keep it "
            "private: don't share the sheet link publicly, and if you use "
            "Google Sheets leave it set to \"Restricted / only you\". It's the "
            "same information you already keep — just organized instead of on "
            "sticky notes.", height=58); r += 2

    head(r, "THE SAMPLE ROWS"); r += 1
    body(r, "The gray italic rows (Sofia Delgado, Marcus Reilly, Priya Nadar "
            "and Aaron Brooks) are examples so you can see how everything works "
            "— including Priya's red CAUTION row (guard dog) and the amber "
            "\"due\" reminders. When you're ready: click the row number on the "
            "left, drag to select the sample rows, right-click, and choose "
            "\"Delete rows\". The formulas and dropdowns stay — only the "
            "example text goes.", height=72); r += 2

    head(r, "ADDING MORE ROWS LATER"); r += 1
    body(r, "Formulas and dropdowns are already loaded far down each tab "
            "(Schedule to row 250, Invoices to row 100, Clients to row 200, "
            "Job Details to row 250). Just keep typing on the next empty line. "
            "If you ever fill a tab completely: select the last row, copy it, "
            "and paste it into the rows below — that carries the formulas down "
            "too.", height=58); r += 2

    head(r, "THE DASHBOARD"); r += 1
    body(r, "100% automatic. Active clients, cleans this month, cleans this "
            "week, cleans today, unpaid balance, revenue this month and your "
            "top clients — it all updates the moment you edit the other tabs.",
         height=44); r += 2

    body(r, "Questions? Message me on Etsy any time — happy to help.  Enjoy! ",
         italic=True); r += 1


# ============================================================================
# build
# ============================================================================

def build():
    wb = Workbook()
    order = ["start", "dash", "clients", "jobs", "schedule", "invoices",
             "services"]
    wb.active.title = S[order[0]]
    for key in order[1:]:
        wb.create_sheet(S[key])

    # defined names for cross-sheet dropdowns (oldest-Excel-safe + GSheets-safe)
    wb.defined_names.add(DefinedName(
        "ClientList",
        attr_text=f"'{S['clients']}'!$A$2:$A${R['clients_max']}"))
    wb.defined_names.add(DefinedName(
        "ServiceList",
        attr_text=f"'{S['services']}'!$A$2:$A${R['services_max']}"))

    build_start_here(wb)
    build_dashboard(wb)
    build_clients(wb)
    build_jobs(wb)
    build_schedule(wb)
    build_invoices(wb)
    build_services(wb)

    wb.calculation.fullCalcOnLoad = True
    out = os.path.join(HERE, CONFIG["output_file"])
    wb.save(out)
    return out


# ============================================================================
# verification
# ============================================================================

# Functions known to behave identically in Excel and Google Sheets.
ALLOWED_FUNCS = {
    "SUMIFS", "SUMIF", "COUNTIF", "COUNTIFS", "COUNTA", "COUNT", "SUM",
    "IF", "IFERROR", "AND", "OR", "NOT", "VLOOKUP", "INDEX", "MATCH",
    "LARGE", "SMALL", "TODAY", "DATE", "YEAR", "MONTH", "DAY", "WEEKDAY",
    "TEXT", "ROUND", "ROW", "COLUMN", "SEARCH", "ISNUMBER", "MIN", "MAX",
}

FUNC_RE = re.compile(r"([A-Z][A-Z0-9\.]*)\(")
SHEET_REF_RE = re.compile(r"'([^']+)'!")


def verify(path):
    problems = []
    wb = load_workbook(path)
    sheet_names = set(wb.sheetnames)
    n_formulas = 0

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                n_formulas += 1
                # strip string literals before analysis
                stripped = re.sub(r'"[^"]*"', '""', v)
                for fn in FUNC_RE.findall(stripped):
                    if fn not in ALLOWED_FUNCS:
                        problems.append(
                            f"{ws.title}!{cell.coordinate}: function {fn} "
                            f"not in dual-compat whitelist")
                for ref in SHEET_REF_RE.findall(stripped):
                    if ref not in sheet_names:
                        problems.append(
                            f"{ws.title}!{cell.coordinate}: reference to "
                            f"missing sheet '{ref}'")
                if stripped.count("(") != stripped.count(")"):
                    problems.append(
                        f"{ws.title}!{cell.coordinate}: unbalanced parens: {v}")

        for dv in ws.data_validations.dataValidation:
            f1 = dv.formula1 or ""
            if f1 in ("ClientList", "ServiceList"):
                continue
            if f1.startswith('"'):
                continue
            m = SHEET_REF_RE.match(f1)
            if m and m.group(1) not in sheet_names:
                problems.append(f"{ws.title} DV: missing sheet in {f1}")

    for name in ("ClientList", "ServiceList"):
        if name not in wb.defined_names:
            problems.append(f"missing defined name {name}")

    print(f"[verify/static] {n_formulas} formula cells scanned")
    return problems


def _lo_profile(tmp):
    """LibreOffice profile that forces recalculation of xlsx on load."""
    profile = os.path.join(tmp, "louser")
    os.makedirs(os.path.join(profile, "user"), exist_ok=True)
    with open(os.path.join(profile, "user",
                           "registrymodifications.xcu"), "w") as f:
        f.write(
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
            'xmlns:xs="http://www.w3.org/2001/XMLSchema">\n'
            '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
            '<prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
            '<value>0</value></prop></item>\n'
            '</oor:items>\n')
    return profile


def verify_recalc(path):
    """Recalculate with LibreOffice (recalc forced via profile setting) and
    look for error values in the computed results."""
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        profile = _lo_profile(tmp)
        cmd = ["soffice", "--headless", "--norestore",
               f"-env:UserInstallation=file://{profile}",
               "--convert-to", "xlsx", "--outdir", tmp, path]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        out = os.path.join(tmp, os.path.basename(path))
        if not os.path.exists(out):
            return [f"LibreOffice recalc failed: {res.stderr.strip()[:300]}"]
        wb = load_workbook(out, data_only=True)
        errs = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!",
                "#NULL!", "Err:")
        n_vals = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if v is None:
                        continue
                    n_vals += 1
                    if isinstance(v, str) and any(e in v for e in errs):
                        problems.append(
                            f"{ws.title}!{cell.coordinate}: computes to {v}")
        print(f"[verify/recalc] LibreOffice recalculated; "
              f"{n_vals} non-empty cells checked for error values")
    return problems


def spot_check(path):
    """Print a few computed values so a human can sanity-check the math."""
    with tempfile.TemporaryDirectory() as tmp:
        profile = _lo_profile(tmp)
        subprocess.run(["soffice", "--headless", "--norestore",
                        f"-env:UserInstallation=file://{profile}",
                        "--convert-to", "xlsx", "--outdir", tmp, path],
                       capture_output=True, text=True, timeout=300)
        wb = load_workbook(os.path.join(tmp, os.path.basename(path)),
                           data_only=True)
        d = wb[S["dash"]]
        print("[spot-check] Dashboard computed values:")
        for label, coord in [("Active Clients", "B6"),
                             ("Cleans This Month", "E6"),
                             ("Cleans This Week", "H6"),
                             ("Cleans Today", "B9"),
                             ("Revenue This Month", "E9"),
                             ("Unpaid Balance", "H9"),
                             ("Top client #1", "C14"),
                             ("Top client #1 revenue", "H14"),
                             ("Top client #2", "C15"),
                             ("Top client #2 revenue", "H15"),
                             ("Top client #3", "C16"),
                             ("Top client #3 revenue", "H16")]:
            print(f"    {label:24s} = {d[coord].value}")
        v = wb[S["schedule"]]
        print("[spot-check] Schedule rates (VLOOKUP from Services):",
              [v[f"D{r}"].value for r in range(2, 8)])
        print("[spot-check] Next-clean-due dates (from cadence):",
              [v[f"G{r}"].value for r in range(2, 8)])
        inv = wb[S["invoices"]]
        print("[spot-check] Invoice totals/balances:",
              [(inv[f"E{r}"].value, inv[f"F{r}"].value) for r in range(2, 5)])


if __name__ == "__main__":
    out = build()
    print(f"built: {out}")
    probs = verify(out)
    probs += verify_recalc(out)
    if probs:
        print("PROBLEMS:")
        for p in probs:
            print("  -", p)
        sys.exit(1)
    spot_check(out)
    print("verification clean.")
