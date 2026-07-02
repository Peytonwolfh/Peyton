#!/usr/bin/env python3
"""
Pet Sitter / Dog Walker Client CRM — Etsy digital product builder.

Generates PetSitter-Client-CRM.xlsx (Google Sheets + Excel compatible):
  * standard formulas only (SUMIFS, COUNTIFS, IF, IFERROR, VLOOKUP, INDEX/MATCH,
    LARGE, TODAY, DATE, WEEKDAY, TEXT) — no VBA, no macros, no Excel-only functions
  * data-validation dropdowns, formula-based conditional formatting
  * parameterized CONFIG so the same script can be re-skinned for other verticals
    (dog groomer, cleaning business, ...) by editing CONFIG only.

Run:  python3 build.py          -> builds + verifies the workbook
"""

import copy
import os
import re
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))

# ============================================================================
# CONFIG — everything vertical-specific lives here.
# To re-skin for another trade (dog groomer, cleaning business, ...):
#   copy this file, edit CONFIG, run.  The builder code below is generic.
# ============================================================================

CONFIG = {
    "output_file": "PetSitter-Client-CRM.xlsx",
    "product_name": "Pet Sitter Client CRM",
    "tagline": "Client, pet, schedule & payment tracker for pet sitters and dog walkers",

    # ---- theme: soft sage green + charcoal -------------------------------
    "theme": {
        "charcoal":      "3B4245",  # header rows
        "charcoal_dark": "2E3436",  # banners
        "sage_dark":     "5C7A5E",  # accents / section heads
        "sage_mid":      "9CB49A",  # tile labels
        "sage_light":    "E3ECE0",  # tile bodies / today-highlight
        "sage_pale":     "F2F6F0",  # page background feel
        "cream":         "FBFAF7",  # input cells
        "calc_fill":     "F1F1EC",  # auto-calculated cells (do not type)
        "red_fill":      "F6D9D5",
        "red_text":      "9C2B23",
        "amber_fill":    "FBEFD0",
        "amber_text":    "8A6116",
        "green_fill":    "DDEBD8",
        "green_text":    "3E6B47",
        "sample_text":   "6D6D6D",  # sample rows: gray italic
        "white":         "FFFFFF",
    },

    # ---- sheet display names (order = tab order) -------------------------
    "sheet_names": {
        "start":    "START HERE",
        "dash":     "Dashboard",
        "clients":  "Clients",
        "pets":     "Pets",
        "visits":   "Visit Schedule",
        "invoices": "Invoices & Payments",
        "services": "Services & Rates",
    },

    # ---- vertical vocabulary (used in labels & START HERE copy) ----------
    "vocab": {
        "client": "Client",
        "unit": "Pet",            # groomer: "Pet"; cleaning biz: "Property"
        "unit_plural": "Pets",
        "visit": "Visit",
        "caution_keyword": "CAUTION",
    },

    # ---- dropdown option lists (inline lists, comma-safe values) ---------
    "lists": {
        "client_status": ["Active", "Inactive", "Waitlist"],
        "contact_method": ["Text", "Call", "Email"],
        "recurring": ["Weekly", "Biweekly", "One-off"],
        "visit_status": ["Booked", "Done", "Cancelled"],
        "paid": ["Yes", "No"],
        "invoice_status": ["Draft", "Sent", "Paid", "Overdue"],
    },

    # ---- editable price list (Services & Rates tab) -----------------------
    # These are NOT sample rows — they're the real starter price list.
    "services": [
        ("Walk 30 min",         22.00, "30 minutes", "One dog; +$8 per extra dog"),
        ("Walk 60 min",         35.00, "60 minutes", "One dog; +$10 per extra dog"),
        ("Drop-in Visit",       28.00, "30 minutes", "Feeding, litter, playtime, meds"),
        ("Overnight Stay",      85.00, "12 hours",   "In client's home, 7pm-7am"),
        ("Boarding (per night)", 65.00, "24 hours",  "In sitter's home"),
    ],

    # ---- column definitions ------------------------------------------------
    # type: input | dropdown:<listkey> | dropdown_range:<defined name> |
    #       date | currency_calc | formula | notes
    "columns": {
        "clients": [
            ("Client Name",               22, "input"),
            ("Phone",                     15, "input"),
            ("Email",                     24, "input"),
            ("Home Address",              28, "input"),
            ("Door / Gate Code",          17, "input"),
            ("Alarm Instructions",        28, "input"),
            ("Key Location",              22, "input"),
            ("Emergency Contact (name & phone)", 26, "input"),
            ("Vet Clinic",                22, "input"),
            ("Vet Phone",                 15, "input"),
            ("Preferred Contact",         16, "dropdown:contact_method"),
            ("Status",                    12, "dropdown:client_status"),
            ("Notes",                     32, "input"),
        ],
        "pets": [
            ("Client",                    22, "dropdown_range:ClientList"),
            ("Pet Name",                  14, "input"),
            ("Species",                   10, "input"),
            ("Breed",                     16, "input"),
            ("Age",                        6, "input"),
            ("Feeding Instructions",      32, "input"),
            ("Medications & Schedule",    30, "input"),
            ("Walk Routine",              28, "input"),
            ("Behavior Notes / Warnings", 34, "input"),
            ("Grooming Notes",            24, "input"),
        ],
        "visits": [
            ("Date",                      12, "date"),
            ("Client",                    22, "dropdown_range:ClientList"),
            ("Pet(s)",                    16, "input"),
            ("Service",                   18, "dropdown_range:ServiceList"),
            ("Recurring?",                12, "dropdown:recurring"),
            ("Rate ($)",                  10, "currency_calc"),
            ("Status",                    11, "dropdown:visit_status"),
            ("Paid?",                      8, "dropdown:paid"),
            ("Notes",                     30, "input"),
        ],
        "invoices": [
            ("Invoice #",                 11, "input"),
            ("Client",                    22, "dropdown_range:ClientList"),
            ("Period Start",              13, "date"),
            ("Period End",                13, "date"),
            ("Invoice Total ($)",         15, "currency_calc"),
            ("Balance Due ($)",           15, "currency_calc"),
            ("Date Sent",                 12, "date"),
            ("Date Paid",                 12, "date"),
            ("Status",                    11, "dropdown:invoice_status"),
            ("Notes",                     30, "input"),
        ],
        "services": [
            ("Service",                   20, "input"),
            ("Rate ($)",                  10, "currency_input"),
            ("Duration",                  12, "input"),
            ("Notes",                     36, "input"),
        ],
    },

    # ---- how deep formulas / dropdowns are pre-loaded ----------------------
    "rows": {
        "clients_max": 200,    # ClientList range + dropdown depth
        "services_max": 30,    # ServiceList range
        "pets_max": 300,
        "visits_max": 250,     # rate formula pre-filled to here
        "invoices_max": 100,   # SUMIFS pre-filled to here
        "top_clients": 5,      # dashboard top-N
        "top_engine": 20,      # first N client rows scanned for top clients
    },

    # ---- sample rows (gray italic, "delete me" note, listed in START HERE) --
    # Dates are TODAY()-relative formulas so the demo conditional formatting
    # (today's visits, unpaid+done) actually lights up for the buyer.
    "sample": {
        "clients": [
            ["Sarah Mitchell", "(555) 201-8843", "sarah.m@example.com",
             "42 Birchwood Ln", "Garage keypad: 4821#",
             "Panel by front door — code 0515, disarm within 60 sec",
             "Lockbox on porch rail — 7734",
             "Tom Mitchell (555) 201-8850", "Lakeview Animal Hospital",
             "(555) 640-2210", "Text", "Active", "SAMPLE ROW — delete me"],
            ["Dana Whitfield", "(555) 318-2276", "danaw@example.com",
             "980 Alder Ct, Unit B", "Gate: 2299 / Door: key only",
             "No alarm", "Under ceramic frog, left of steps",
             "Priya Nair (sister) (555) 318-9901", "Eastside Vet Clinic",
             "(555) 772-0184", "Email", "Active", "SAMPLE ROW — delete me"],
            ["Marcus Lee", "(555) 467-9032", "marcus.lee@example.com",
             "17 Harborview Dr", "Smart lock — code texted weekly",
             "Alarm app on my phone (shared access)", "N/A — smart lock",
             "Jen Lee (555) 467-9033", "Harbor Animal Hospital",
             "(555) 903-3341", "Call", "Waitlist", "SAMPLE ROW — delete me"],
        ],
        "pets": [
            ["Sarah Mitchell", "Biscuit", "Dog", "Golden Retriever", "4",
             "1 cup kibble 7am & 6pm; treats after walks OK",
             "None",
             "30 min neighborhood loop; pulls hard at squirrels",
             "Friendly with all dogs & people",
             "Brush weekly — SAMPLE ROW, delete me"],
            ["Sarah Mitchell", "Clover", "Cat", "Domestic Shorthair", "7",
             "1/2 can wet food 8am; dry food always out",
             "Thyroid pill 8am — hide in Churu treat",
             "Indoor only",
             "Shy — hides under guest bed from strangers",
             "SAMPLE ROW — delete me"],
            ["Dana Whitfield", "Rocko", "Dog", "Boxer mix", "2",
             "2 cups kibble 6pm only; no chicken (allergy)",
             "None",
             "60 min run or fetch at Alder Park",
             "CAUTION: reactive to bikes & skateboards — short leash near paths",
             "SAMPLE ROW — delete me"],
            ["Marcus Lee", "Pepper", "Dog", "Mini Schnauzer", "9",
             "1/2 cup senior kibble 7am & 5pm",
             "Rimadyl 25mg with dinner",
             "20 min gentle walk, no stairs",
             "Barks at mail carriers — all noise, no bite",
             "Grooming every 6 weeks — SAMPLE ROW, delete me"],
        ],
        # date offset (days from TODAY, or a literal formula string),
        # client, pets, service, recurring, status, paid, note
        "visits": [
            [0,   "Sarah Mitchell", "Biscuit + Clover", "Drop-in Visit",
             "Weekly",  "Booked", "No",  "SAMPLE — today's visits glow green"],
            [-3,  "Dana Whitfield", "Rocko", "Walk 60 min",
             "Weekly",  "Done",   "No",  "SAMPLE — done + unpaid turns red"],
            [-7,  "Sarah Mitchell", "Biscuit", "Walk 30 min",
             "Weekly",  "Done",   "Yes", "SAMPLE ROW — delete me"],
            # anchored to the 1st of the current month so the "Revenue This
            # Month" dashboard tile always has something to show
            ["=DATE(YEAR(TODAY()),MONTH(TODAY()),1)",
             "Dana Whitfield", "Rocko", "Walk 60 min",
             "Weekly",  "Done",   "Yes", "SAMPLE ROW — delete me"],
            [-10, "Marcus Lee", "Pepper", "Overnight Stay",
             "One-off", "Done",   "No",  "SAMPLE ROW — delete me"],
            [2,   "Marcus Lee", "Pepper", "Walk 30 min",
             "Biweekly", "Booked", "No", "SAMPLE ROW — delete me"],
        ],
        # inv#, client, period_start_offset, period_end_offset,
        # sent_offset (None = blank), paid_offset (None = blank), status, note
        "invoices": [
            ["INV-1001", "Sarah Mitchell", -14, -5, -4, -2, "Paid",
             "SAMPLE ROW — delete me"],
            ["INV-1002", "Dana Whitfield", -6, 0, 0, None, "Sent",
             "SAMPLE ROW — delete me"],
            ["INV-1003", "Marcus Lee", -20, -8, -18, None, "Overdue",
             "SAMPLE ROW — delete me"],
        ],
    },
}

T = CONFIG["theme"]
S = CONFIG["sheet_names"]
R = CONFIG["rows"]

# ============================================================================
# style helpers
# ============================================================================

def fill(hex_):
    return PatternFill("solid", fgColor=hex_)

THIN = Side(style="thin", color="D8DDD6")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_FONT = Font(name="Calibri", bold=True, color=T["white"], size=11)
HEADER_FILL = fill(T["charcoal"])
SAMPLE_FONT = Font(name="Calibri", italic=True, color=T["sample_text"], size=11)
BODY_FONT = Font(name="Calibri", size=11)

CURRENCY_FMT = '"$"#,##0.00'
DATE_FMT = "mm/dd/yyyy"


def style_header_row(ws, n_cols, row=1, height=26):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.row_dimensions[row].height = height


def set_widths(ws, cols):
    for i, (_, width, _) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def apply_col_formats(ws, cols, first_row, last_row):
    """Number formats + calc-cell shading for whole data region."""
    for i, (_, _, ctype) in enumerate(cols, start=1):
        for r in range(first_row, last_row + 1):
            cell = ws.cell(row=r, column=i)
            if ctype in ("currency_calc", "currency_input"):
                cell.number_format = CURRENCY_FMT
            elif ctype == "date":
                cell.number_format = DATE_FMT
            if ctype == "currency_calc":
                cell.fill = fill(T["calc_fill"])


def mark_sample_rows(ws, n_cols, first_row, n_rows):
    for r in range(first_row, first_row + n_rows):
        for c in range(1, n_cols + 1):
            ws.cell(row=r, column=c).font = SAMPLE_FONT


def add_dv(ws, dv_type, formula1, cell_range):
    dv = DataValidation(type=dv_type, formula1=formula1, allow_blank=True,
                        showErrorMessage=False)
    ws.add_data_validation(dv)
    dv.add(cell_range)
    return dv


def inline_list(key):
    return '"' + ",".join(CONFIG["lists"][key]) + '"'


def build_columns(ws, sheet_key, max_row):
    """Headers, widths, dropdown DV for a standard data sheet."""
    cols = CONFIG["columns"][sheet_key]
    for i, (label, _, _) in enumerate(cols, start=1):
        ws.cell(row=1, column=i, value=label)
    style_header_row(ws, len(cols))
    set_widths(ws, cols)
    ws.freeze_panes = "A2"
    for i, (_, _, ctype) in enumerate(cols, start=1):
        col = get_column_letter(i)
        rng = f"{col}2:{col}{max_row}"
        if ctype.startswith("dropdown_range:"):
            add_dv(ws, "list", ctype.split(":", 1)[1], rng)
        elif ctype.startswith("dropdown:"):
            add_dv(ws, "list", inline_list(ctype.split(":", 1)[1]), rng)
    apply_col_formats(ws, cols, 2, max_row)
    return cols


# ============================================================================
# sheet builders
# ============================================================================

def build_services(wb):
    ws = wb[S["services"]]
    cols = build_columns(ws, "services", R["services_max"])
    for r, row in enumerate(CONFIG["services"], start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    # gentle banding on the real price rows
    for r in range(2, len(CONFIG["services"]) + 2):
        for c in range(1, len(cols) + 1):
            ws.cell(row=r, column=c).border = BORDER_ALL
    note = ws.cell(row=R["services_max"] + 2, column=1,
                   value="Edit names and rates freely — the Visit Schedule "
                         "looks prices up from this list automatically. "
                         "Service names here must match the dropdown picks "
                         "exactly (they will, if you always pick from the "
                         "dropdown).")
    note.font = Font(italic=True, color=T["sample_text"], size=10)
    ws.merge_cells(start_row=R["services_max"] + 2, start_column=1,
                   end_row=R["services_max"] + 2, end_column=4)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_properties.tabColor = T["sage_dark"]


def build_clients(wb):
    ws = wb[S["clients"]]
    build_columns(ws, "clients", R["clients_max"])
    samples = CONFIG["sample"]["clients"]
    for r, row in enumerate(samples, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    mark_sample_rows(ws, len(CONFIG["columns"]["clients"]), 2, len(samples))
    # status column CF: Active = green tint, Waitlist = amber
    status_col = get_column_letter(
        [l for l, _, _ in CONFIG["columns"]["clients"]].index("Status") + 1)
    rng = f"A2:{get_column_letter(len(CONFIG['columns']['clients']))}{R['clients_max']}"
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{R['clients_max']}",
        FormulaRule(formula=[f'${status_col}2="Active"'],
                    fill=fill(T["green_fill"]),
                    font=Font(color=T["green_text"], bold=True)))
    ws.conditional_formatting.add(
        f"{status_col}2:{status_col}{R['clients_max']}",
        FormulaRule(formula=[f'${status_col}2="Waitlist"'],
                    fill=fill(T["amber_fill"]),
                    font=Font(color=T["amber_text"], bold=True)))
    ws.sheet_properties.tabColor = T["sage_mid"]


def build_pets(wb):
    ws = wb[S["pets"]]
    cols = build_columns(ws, "pets", R["pets_max"])
    samples = CONFIG["sample"]["pets"]
    for r, row in enumerate(samples, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    # CAUTION keyword anywhere in the behavior column -> whole row red
    beh_col = get_column_letter(
        [l for l, _, _ in cols].index("Behavior Notes / Warnings") + 1)
    last_col = get_column_letter(len(cols))
    kw = CONFIG["vocab"]["caution_keyword"]
    ws.conditional_formatting.add(
        f"A2:{last_col}{R['pets_max']}",
        FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{kw}",${beh_col}2))'],
            fill=fill(T["red_fill"]),
            font=Font(color=T["red_text"]),
            stopIfTrue=True))
    ws.sheet_properties.tabColor = T["sage_mid"]


def build_visits(wb):
    ws = wb[S["visits"]]
    cols = build_columns(ws, "visits", R["visits_max"])
    labels = [l for l, _, _ in cols]
    c_date = get_column_letter(labels.index("Date") + 1)          # A
    c_service = get_column_letter(labels.index("Service") + 1)    # D
    c_rate = get_column_letter(labels.index("Rate ($)") + 1)      # F
    c_status = get_column_letter(labels.index("Status") + 1)      # G
    c_paid = get_column_letter(labels.index("Paid?") + 1)         # H
    last_col = get_column_letter(len(cols))
    svc_sheet = S["services"]

    # pre-load the rate lookup all the way down
    for r in range(2, R["visits_max"] + 1):
        ws.cell(row=r, column=labels.index("Rate ($)") + 1).value = (
            f'=IF(${c_service}{r}="","",'
            f"IFERROR(VLOOKUP(${c_service}{r},"
            f"'{svc_sheet}'!$A$2:$B${R['services_max']},2,FALSE),\"\"))"
        )

    samples = CONFIG["sample"]["visits"]
    for r, (off, client, pets, service, recur, status, paid, note) in \
            enumerate(samples, start=2):
        if isinstance(off, str):
            date_formula = off
        else:
            date_formula = (f"=TODAY()+{off}" if off >= 0
                            else f"=TODAY()-{abs(off)}")
        ws.cell(row=r, column=1, value=date_formula).number_format = DATE_FMT
        ws.cell(row=r, column=2, value=client)
        ws.cell(row=r, column=3, value=pets)
        ws.cell(row=r, column=4, value=service)
        ws.cell(row=r, column=5, value=recur)
        ws.cell(row=r, column=7, value=status)
        ws.cell(row=r, column=8, value=paid)
        ws.cell(row=r, column=9, value=note)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    # keep number formats after sample write
    apply_col_formats(ws, cols, 2, R["visits_max"])

    # CF 1 (wins): done + not paid -> red row
    ws.conditional_formatting.add(
        f"A2:{last_col}{R['visits_max']}",
        FormulaRule(
            formula=[f'AND(${c_status}2="Done",${c_paid}2<>"Yes")'],
            fill=fill(T["red_fill"]), font=Font(color=T["red_text"]),
            stopIfTrue=True))
    # CF 2: today's visits -> sage highlight
    ws.conditional_formatting.add(
        f"A2:{last_col}{R['visits_max']}",
        FormulaRule(formula=[f"${c_date}2=TODAY()"],
                    fill=fill(T["sage_light"]),
                    font=Font(color=T["charcoal_dark"], bold=True),
                    stopIfTrue=True))
    ws.sheet_properties.tabColor = T["sage_mid"]


def build_invoices(wb):
    ws = wb[S["invoices"]]
    cols = build_columns(ws, "invoices", R["invoices_max"])
    labels = [l for l, _, _ in cols]
    c_client, c_start, c_end = "B", "C", "D"
    c_status = get_column_letter(labels.index("Status") + 1)
    last_col = get_column_letter(len(cols))
    vs = S["visits"]

    for r in range(2, R["invoices_max"] + 1):
        common = (f"'{vs}'!$F:$F,'{vs}'!$B:$B,${c_client}{r},"
                  f"'{vs}'!$G:$G,\"Done\","
                  f"'{vs}'!$A:$A,\">=\"&${c_start}{r},"
                  f"'{vs}'!$A:$A,\"<=\"&${c_end}{r}")
        guard = (f'OR(${c_client}{r}="",${c_start}{r}="",${c_end}{r}="")')
        ws.cell(row=r, column=labels.index("Invoice Total ($)") + 1).value = \
            f'=IF({guard},"",SUMIFS({common}))'
        ws.cell(row=r, column=labels.index("Balance Due ($)") + 1).value = \
            f'=IF({guard},"",SUMIFS({common},\'{vs}\'!$H:$H,"<>Yes"))'

    def off_formula(off):
        if off is None:
            return None
        return f"=TODAY()+{off}" if off >= 0 else f"=TODAY()-{abs(off)}"

    samples = CONFIG["sample"]["invoices"]
    for r, (num, client, ps, pe, sent, paid, status, note) in \
            enumerate(samples, start=2):
        ws.cell(row=r, column=1, value=num)
        ws.cell(row=r, column=2, value=client)
        ws.cell(row=r, column=3, value=off_formula(ps))
        ws.cell(row=r, column=4, value=off_formula(pe))
        if sent is not None:
            ws.cell(row=r, column=7, value=off_formula(sent))
        if paid is not None:
            ws.cell(row=r, column=8, value=off_formula(paid))
        ws.cell(row=r, column=9, value=status)
        ws.cell(row=r, column=10, value=note)
    mark_sample_rows(ws, len(cols), 2, len(samples))
    apply_col_formats(ws, cols, 2, R["invoices_max"])

    # status CF on whole row
    rng = f"A2:{last_col}{R['invoices_max']}"
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'${c_status}2="Overdue"'], fill=fill(T["red_fill"]),
        font=Font(color=T["red_text"]), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'${c_status}2="Paid"'], fill=fill(T["green_fill"]),
        font=Font(color=T["green_text"]), stopIfTrue=True))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'${c_status}2="Sent"'], fill=fill(T["amber_fill"]),
        font=Font(color=T["amber_text"]), stopIfTrue=True))
    ws.sheet_properties.tabColor = T["sage_mid"]


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

def _tile(ws, col, label, formula, num_fmt, label_row, value_row):
    """Two merged 2-wide cells stacked: label on top, big value below."""
    c1, c2 = col, get_column_letter(column_index_from_string(col) + 1)
    ws.merge_cells(f"{c1}{label_row}:{c2}{label_row}")
    ws.merge_cells(f"{c1}{value_row}:{c2}{value_row}")
    lab = ws[f"{c1}{label_row}"]
    lab.value = label.upper()
    lab.font = Font(bold=True, size=9, color=T["white"])
    lab.fill = fill(T["sage_dark"])
    lab.alignment = Alignment(horizontal="center", vertical="center")
    val = ws[f"{c1}{value_row}"]
    val.value = formula
    val.font = Font(bold=True, size=22, color=T["charcoal_dark"])
    val.fill = fill(T["sage_light"])
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.number_format = num_fmt
    for cc in (c1, c2):
        ws[f"{cc}{label_row}"].fill = fill(T["sage_dark"])
        ws[f"{cc}{value_row}"].fill = fill(T["sage_light"])


def build_dashboard(wb):
    ws = wb[S["dash"]]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = T["charcoal_dark"]

    cl, vs = S["clients"], S["visits"]
    cl_labels = [l for l, _, _ in CONFIG["columns"]["clients"]]
    cl_status = get_column_letter(cl_labels.index("Status") + 1)  # L

    # widths: A spacer, B..I tiles, then helper cols
    ws.column_dimensions["A"].width = 2
    for c in "BCEFHI":
        ws.column_dimensions[c].width = 13
    for c in "DG":
        ws.column_dimensions[c].width = 2
    ws.column_dimensions["J"].width = 2
    ws.column_dimensions["O"].width = 12
    ws.column_dimensions["P"].width = 18

    # title band
    ws.merge_cells("B2:I2")
    t = ws["B2"]
    t.value = CONFIG["product_name"].upper() + "  —  DASHBOARD"
    t.font = Font(bold=True, size=16, color=T["white"])
    t.fill = fill(T["charcoal_dark"])
    t.alignment = Alignment(horizontal="center", vertical="center")
    for c in "CDEFGHI":
        ws[f"{c}2"].fill = fill(T["charcoal_dark"])
    ws.row_dimensions[2].height = 30

    ws["B3"] = '="Everything on this page updates itself — nothing to type here.  Today: "&TEXT(TODAY(),"mmm d, yyyy")'
    ws["B3"].font = Font(italic=True, size=10, color=T["sample_text"])

    week_start = "(TODAY()-WEEKDAY(TODAY(),2)+1)"
    tiles = [
        ("B", "Active Clients",
         f"=COUNTIF('{cl}'!${cl_status}:${cl_status},\"Active\")", "0"),
        ("E", "Pets Under Care",
         f"=COUNTA('{S['pets']}'!$B$2:$B${R['pets_max']})", "0"),
        ("H", "Visits This Week",
         f"=COUNTIFS('{vs}'!$A:$A,\">=\"&{week_start},"
         f"'{vs}'!$A:$A,\"<=\"&{week_start}+6,"
         f"'{vs}'!$G:$G,\"<>Cancelled\")", "0"),
    ]
    tiles2 = [
        ("B", "Visits Today",
         f"=COUNTIFS('{vs}'!$A:$A,TODAY(),'{vs}'!$G:$G,\"<>Cancelled\")", "0"),
        ("E", "Revenue This Month",
         f"=SUMIFS('{vs}'!$F:$F,'{vs}'!$G:$G,\"Done\","
         f"'{vs}'!$A:$A,\">=\"&DATE(YEAR(TODAY()),MONTH(TODAY()),1),"
         f"'{vs}'!$A:$A,\"<\"&DATE(YEAR(TODAY()),MONTH(TODAY())+1,1))",
         CURRENCY_FMT),
        ("H", "Unpaid Balance",
         f"=SUMIFS('{vs}'!$F:$F,'{vs}'!$G:$G,\"Done\",'{vs}'!$H:$H,\"<>Yes\")",
         CURRENCY_FMT),
    ]
    for col, label, formula, fmt in tiles:
        _tile(ws, col, label, formula, fmt, 5, 6)
    for col, label, formula, fmt in tiles2:
        _tile(ws, col, label, formula, fmt, 8, 9)
    ws.row_dimensions[6].height = 40
    ws.row_dimensions[9].height = 40

    # ---- top clients table -------------------------------------------------
    ws.merge_cells("B12:I12")
    h = ws["B12"]
    h.value = "TOP CLIENTS  (completed visits, all time)"
    h.font = Font(bold=True, size=12, color=T["white"])
    h.fill = fill(T["charcoal"])
    h.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for c in "CDEFGHI":
        ws[f"{c}12"].fill = fill(T["charcoal"])
    ws.row_dimensions[12].height = 22

    hdr = {"B": "Rank", "C": "Client", "H": "Revenue"}
    ws.merge_cells("C13:G13")
    for c, v in hdr.items():
        cell = ws[f"{c}13"]
        cell.value = v
        cell.font = Font(bold=True, size=10, color=T["sage_dark"])
        cell.fill = fill(T["sage_pale"])
    ws.merge_cells("H13:I13")
    ws["H13"].fill = fill(T["sage_pale"])

    eng_top, eng_n = 14, R["top_engine"]           # engine rows 14..33
    eng_rng = f"$O${eng_top}:$O${eng_top + eng_n - 1}"
    name_rng = f"$P${eng_top}:$P${eng_top + eng_n - 1}"
    for k in range(1, R["top_clients"] + 1):
        row = 13 + k
        ws[f"B{row}"] = k
        ws[f"B{row}"].alignment = Alignment(horizontal="center")
        ws.merge_cells(f"C{row}:G{row}")
        ws[f"C{row}"] = (f"=IFERROR(INDEX({name_rng},"
                         f"MATCH(LARGE({eng_rng},{k}),{eng_rng},0)),\"—\")")
        ws.merge_cells(f"H{row}:I{row}")
        ws[f"H{row}"] = f'=IFERROR(ROUND(LARGE({eng_rng},{k}),2),"")'
        ws[f"H{row}"].number_format = CURRENCY_FMT
        for c in "BCDEFGHI":
            ws[f"{c}{row}"].fill = fill(
                T["sage_pale"] if k % 2 else T["white"])
        ws[f"C{row}"].font = Font(size=11, color=T["charcoal_dark"])
        ws[f"H{row}"].font = Font(size=11, bold=True, color=T["sage_dark"])

    # ---- hidden-ish calculation engine (cols O:P) ---------------------------
    ws["O13"] = "engine ↓"
    ws["P13"] = "do not edit"
    for c in ("O13", "P13"):
        ws[c].font = Font(size=8, color=T["sample_text"], italic=True)
    for i in range(eng_n):
        r = eng_top + i
        cl_row = 2 + i
        ws[f"O{r}"] = (f"=IF('{cl}'!A{cl_row}=\"\",\"\","
                       f"SUMIFS('{vs}'!$F:$F,'{vs}'!$B:$B,'{cl}'!A{cl_row},"
                       f"'{vs}'!$G:$G,\"Done\")+ROW()/10000000)")
        ws[f"P{r}"] = f"=IF($O{r}=\"\",\"\",'{cl}'!A{cl_row})"
        for c in ("O", "P"):
            ws[f"{c}{r}"].font = Font(size=8, color=T["sample_text"])
    ws.column_dimensions["O"].hidden = True
    ws.column_dimensions["P"].hidden = True

    note = ws["B21"]
    note.value = ("Top-client list scans your first "
                  f"{eng_n} client rows. Hidden columns O-P do the math — "
                  "leave them alone (Format > Unhide if you're curious).")
    note.font = Font(italic=True, size=9, color=T["sample_text"])
    ws.merge_cells("B21:I21")


# ---------------------------------------------------------------------------
# START HERE
# ---------------------------------------------------------------------------

def build_start_here(wb):
    ws = wb[S["start"]]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = T["sage_dark"]
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 100
    ws.column_dimensions["C"].width = 3

    def head(row, text, big=False):
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = Font(bold=True, size=16 if big else 12,
                         color=T["white"] if big else T["white"])
        cell.fill = fill(T["charcoal_dark"] if big else T["sage_dark"])
        cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 34 if big else 24

    def body(row, text, italic=False, height=None):
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = Font(size=11, italic=italic, color=T["charcoal_dark"])
        cell.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        cell.fill = fill(T["cream"])
        if height:
            ws.row_dimensions[row].height = height
        return cell

    r = 2
    head(r, f"WELCOME TO YOUR {CONFIG['product_name'].upper()}", big=True); r += 1
    body(r, "Made for pet sitters and dog walkers. Everything your business "
            "needs — clients, door codes, vet info, medications, schedule and "
            "money — in one place. No apps, no subscriptions, nothing to "
            "install.", height=44); r += 2

    head(r, "STEP 1 — MAKE YOUR OWN COPY (do this first!)"); r += 1
    body(r, "USING GOOGLE SHEETS (free):  1. Go to sheets.google.com and sign "
            "in.  2. Click the folder icon (Open file picker), choose "
            "\"Upload\", and drop this file in.  3. It opens right in Google "
            "Sheets — then click File > Save as Google Sheets so it lives in "
            "your Drive forever.", height=58); r += 1
    body(r, "USING EXCEL:  just open the file and File > Save As with your "
            "business name. Done.", height=28); r += 2

    head(r, "STEP 2 — KNOW YOUR COLORS"); r += 1
    body(r, "WHITE cells = yours. Type anything.", height=20); r += 1
    c = body(r, "LIGHT GRAY cells = the spreadsheet's. They calculate "
                "themselves (rates, totals, the whole Dashboard). "
                "Don't type in them — if you ever do by accident, just press "
                "Ctrl+Z (Cmd+Z on Mac).", height=44)
    c.fill = fill(T["calc_fill"]); r += 1
    body(r, "Cells with a little arrow = dropdowns. Click the arrow and pick "
            "— never type these by hand, so names always match everywhere.",
         height=30); r += 2

    head(r, "STEP 3 — SET UP IN THIS ORDER"); r += 1
    body(r, "1. SERVICES & RATES — change the service names and prices to "
            "yours. The schedule fills in prices from this list "
            "automatically.", height=30); r += 1
    body(r, "2. CLIENTS — one row per household. Door codes, alarm notes, "
            "key location, vet and emergency contact all live here, so "
            "everything you need mid-visit is one tap away on your phone.",
         height=44); r += 1
    body(r, "3. PETS — one row per pet. Pick the owner from the dropdown. "
            "Tip: type the word CAUTION anywhere in \"Behavior Notes\" and "
            "the whole row turns red — impossible to miss before a visit.",
         height=44); r += 1
    body(r, "4. VISIT SCHEDULE — one row per visit. Pick client and service "
            "from dropdowns; the rate appears by itself. Mark Status = Done "
            "when finished and Paid? = Yes when the money arrives. Today's "
            "visits glow green; finished-but-unpaid visits turn red.",
         height=58); r += 1
    body(r, "5. INVOICES & PAYMENTS — type an invoice number, pick the "
            "client, and type the period start and end dates. The totals add "
            "themselves up from the Visit Schedule (completed visits in that "
            "period; Balance Due counts only the unpaid ones).", height=44); r += 2

    head(r, "THE SAMPLE ROWS"); r += 1
    body(r, "The gray italic rows (Sarah Mitchell, Dana Whitfield, Marcus "
            "Lee...) are examples so you can see how everything works. When "
            "you're ready: click the row number on the left, drag to select "
            "the sample rows, right-click, and choose \"Delete rows\". The "
            "formulas and dropdowns stay — only the example text goes.",
         height=58); r += 2

    head(r, "ADDING MORE ROWS LATER"); r += 1
    body(r, "Formulas and dropdowns are already loaded far down each tab "
            "(Visit Schedule to row 250, Invoices to row 100, Clients to row "
            "200). Just keep typing on the next empty line. If you ever fill "
            "a tab completely: select the last row, copy it, and paste it "
            "into the rows below — that carries the formulas down too.",
         height=58); r += 2

    head(r, "THE DASHBOARD"); r += 1
    body(r, "100% automatic. Active clients, pets under care, this week's "
            "visits, unpaid balance, this month's revenue and your top "
            "clients — it all updates the moment you edit the other tabs.",
         height=44); r += 2

    body(r, "Questions? Message me on Etsy any time — happy to help.  Enjoy! ",
         italic=True); r += 1


# ============================================================================
# build
# ============================================================================

def build():
    wb = Workbook()
    order = ["start", "dash", "clients", "pets", "visits", "invoices",
             "services"]
    wb.active.title = S[order[0]]
    for key in order[1:]:
        wb.create_sheet(S[key])

    # defined names for cross-sheet dropdowns (oldest-Excel-safe + GSheets-safe)
    wb.defined_names.add(DefinedName(
        "ClientList",
        attr_text=f"'{S['clients']}'!$A$2:$A${R['clients_max']}"))
    wb.defined_names.add(DefinedName(
        "ServiceList",
        attr_text=f"'{S['services']}'!$A$2:$A${R['services_max']}"))

    build_start_here(wb)
    build_dashboard(wb)
    build_clients(wb)
    build_pets(wb)
    build_visits(wb)
    build_invoices(wb)
    build_services(wb)

    wb.calculation.fullCalcOnLoad = True
    out = os.path.join(HERE, CONFIG["output_file"])
    wb.save(out)
    return out


# ============================================================================
# verification
# ============================================================================

# Functions known to behave identically in Excel and Google Sheets.
ALLOWED_FUNCS = {
    "SUMIFS", "SUMIF", "COUNTIF", "COUNTIFS", "COUNTA", "COUNT", "SUM",
    "IF", "IFERROR", "AND", "OR", "NOT", "VLOOKUP", "INDEX", "MATCH",
    "LARGE", "SMALL", "TODAY", "DATE", "YEAR", "MONTH", "DAY", "WEEKDAY",
    "TEXT", "ROUND", "ROW", "COLUMN", "SEARCH", "ISNUMBER", "MIN", "MAX",
}

FUNC_RE = re.compile(r"([A-Z][A-Z0-9\.]*)\(")
SHEET_REF_RE = re.compile(r"'([^']+)'!")


def verify(path):
    problems = []
    wb = load_workbook(path)
    sheet_names = set(wb.sheetnames)
    n_formulas = 0

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                n_formulas += 1
                # strip string literals before analysis
                stripped = re.sub(r'"[^"]*"', '""', v)
                for fn in FUNC_RE.findall(stripped):
                    if fn not in ALLOWED_FUNCS:
                        problems.append(
                            f"{ws.title}!{cell.coordinate}: function {fn} "
                            f"not in dual-compat whitelist")
                for ref in SHEET_REF_RE.findall(stripped):
                    if ref not in sheet_names:
                        problems.append(
                            f"{ws.title}!{cell.coordinate}: reference to "
                            f"missing sheet '{ref}'")
                if stripped.count("(") != stripped.count(")"):
                    problems.append(
                        f"{ws.title}!{cell.coordinate}: unbalanced parens: {v}")

        for dv in ws.data_validations.dataValidation:
            f1 = dv.formula1 or ""
            if f1 in ("ClientList", "ServiceList"):
                continue
            if f1.startswith('"'):
                continue
            m = SHEET_REF_RE.match(f1)
            if m and m.group(1) not in sheet_names:
                problems.append(f"{ws.title} DV: missing sheet in {f1}")

    for name in ("ClientList", "ServiceList"):
        if name not in wb.defined_names:
            problems.append(f"missing defined name {name}")

    print(f"[verify/static] {n_formulas} formula cells scanned")
    return problems


def verify_recalc(path):
    """Recalculate with LibreOffice (recalc forced via profile setting) and
    look for error values in the computed results."""
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        profile = os.path.join(tmp, "louser")
        os.makedirs(os.path.join(profile, "user"), exist_ok=True)
        # force "always recalculate on load" for xlsx
        with open(os.path.join(profile, "user",
                               "registrymodifications.xcu"), "w") as f:
            f.write(
                '<?xml version="1.0" encoding="UTF-8"?>\n'
                '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
                'xmlns:xs="http://www.w3.org/2001/XMLSchema">\n'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
                '<prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
                '<value>0</value></prop></item>\n'
                '</oor:items>\n')
        cmd = ["soffice", "--headless", "--norestore",
               f"-env:UserInstallation=file://{profile}",
               "--convert-to", "xlsx", "--outdir", tmp, path]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        out = os.path.join(tmp, os.path.basename(path))
        if not os.path.exists(out):
            return [f"LibreOffice recalc failed: {res.stderr.strip()[:300]}"]
        wb = load_workbook(out, data_only=True)
        errs = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!",
                "#NULL!", "Err:")
        n_vals = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if v is None:
                        continue
                    n_vals += 1
                    if isinstance(v, str) and any(e in v for e in errs):
                        problems.append(
                            f"{ws.title}!{cell.coordinate}: computes to {v}")
        print(f"[verify/recalc] LibreOffice recalculated; "
              f"{n_vals} non-empty cells checked for error values")
    return problems


def spot_check(path):
    """Print a few computed values so a human can sanity-check the math."""
    with tempfile.TemporaryDirectory() as tmp:
        profile = os.path.join(tmp, "louser")
        os.makedirs(os.path.join(profile, "user"), exist_ok=True)
        with open(os.path.join(profile, "user",
                               "registrymodifications.xcu"), "w") as f:
            f.write(
                '<?xml version="1.0" encoding="UTF-8"?>\n'
                '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
                'xmlns:xs="http://www.w3.org/2001/XMLSchema">\n'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
                '<prop oor:name="OOXMLRecalcMode" oor:op="fuse">'
                '<value>0</value></prop></item>\n'
                '</oor:items>\n')
        subprocess.run(["soffice", "--headless", "--norestore",
                        f"-env:UserInstallation=file://{profile}",
                        "--convert-to", "xlsx", "--outdir", tmp, path],
                       capture_output=True, text=True, timeout=300)
        wb = load_workbook(os.path.join(tmp, os.path.basename(path)),
                           data_only=True)
        d = wb[S["dash"]]
        print("[spot-check] Dashboard computed values:")
        for label, coord in [("Active Clients", "B6"),
                             ("Pets Under Care", "E6"),
                             ("Visits This Week", "H6"),
                             ("Visits Today", "B9"),
                             ("Revenue This Month", "E9"),
                             ("Unpaid Balance", "H9"),
                             ("Top client #1", "C14"),
                             ("Top client #1 revenue", "H14")]:
            print(f"    {label:24s} = {d[coord].value}")
        v = wb[S["visits"]]
        print("[spot-check] Visit rates (VLOOKUP):",
              [v[f"F{r}"].value for r in range(2, 7)])
        inv = wb[S["invoices"]]
        print("[spot-check] Invoice totals/balances:",
              [(inv[f"E{r}"].value, inv[f"F{r}"].value) for r in range(2, 5)])


if __name__ == "__main__":
    out = build()
    print(f"built: {out}")
    probs = verify(out)
    probs += verify_recalc(out)
    if probs:
        print("PROBLEMS:")
        for p in probs:
            print("  -", p)
        sys.exit(1)
    spot_check(out)
    print("verification clean.")
